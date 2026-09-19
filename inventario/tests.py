import tempfile
from datetime import date
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

from .models import (
    ConteoDetalle,
    ConteoFisico,
    CostoAlmacenamiento,
    Movimiento,
    Pedido,
    PedidoDetalle,
    Producto,
)


class MovimientoStockTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username='tester', password='pass12345')
        self.producto = Producto.objects.create(
            codigo='PIN-001',
            nombre='Pintura latex',
            precio_venta='50.00',
            costo_compra='30.00',
            stock_actual=10,
            stock_minimo=2,
        )

    def _crear_movimiento(self, tipo, cantidad):
        return Movimiento.objects.create(
            producto=self.producto,
            tipo=tipo,
            cantidad=cantidad,
            fecha=timezone.now(),
            usuario=self.usuario,
        )

    def test_ingreso_suma_stock(self):
        self._crear_movimiento(Movimiento.Tipo.INGRESO, 5)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, 15)

    def test_salida_resta_stock(self):
        self._crear_movimiento(Movimiento.Tipo.SALIDA, 4)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, 6)

    def test_salida_mayor_al_stock_falla(self):
        with self.assertRaises(ValidationError):
            self._crear_movimiento(Movimiento.Tipo.SALIDA, 100)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, 10)

    def test_borrar_movimiento_revierte_stock(self):
        movimiento = self._crear_movimiento(Movimiento.Tipo.INGRESO, 7)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, 17)

        movimiento.delete()
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, 10)


def _construir_excel_minimo(ruta):
    """Arma un Excel de 3 filas por hoja que reproduce la estructura real de
    las fichas de registro (mismas columnas y encabezados que
    datos/fichas_pretest_agosto_2026.xlsx, incluidas las que el comando no
    lee pero sí exige o que existen en el archivo real), para probar el
    comando cargar_datos sin depender del archivo de la tesis."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    ei = libro.create_sheet('1_EI')
    for _ in range(10):
        ei.append([])
    ei.append([
        'N.°', 'Fecha de corte', 'Código', 'Producto y presentación',
        'Stock según kardex', 'Stock según conteo físico', 'Diferencia',
        '¿Registro correcto? (1/0)',
    ])
    ei.append([1, date(2026, 9, 10), 'PIN-001', 'CPP Pato — Blanco — 1/4 galón', 10, 10, 0, 1])
    ei.append([2, date(2026, 9, 10), 'PIN-002', 'Innova — Celeste — balde 4 galones', 9, 3, -6, 0])
    ei.append([3, date(2026, 9, 10), 'PIN-003', 'Producto sin formato válido', 8, 8, 0, 1])

    # 2_NS trae "Código del producto" y "Motivo" en el archivo real
    # (fichas_pretest_agosto_2026.xlsx); "Motivo" no lo lee el comando, pero
    # se incluye igual para que la estructura de columnas coincida con la
    # real. El pedido 3 deja el código en blanco a propósito, para seguir
    # probando el caso (real, aunque raro) de un producto sin código.
    ns = libro.create_sheet('2_NS')
    for _ in range(10):
        ns.append([])
    ns.append([
        'N.°', 'Fecha del pedido', 'Cliente (código)', 'Código del producto',
        'Producto y presentación', 'Cant. solicitada', 'Cant. atendida',
        'Fecha requerida', 'Fecha de atención', 'Motivo', '¿A tiempo? (1/0)',
    ])
    ns.append([
        1, date(2026, 10, 1), 'C-001', 'PIN-001', 'CPP Pato — Blanco — 1/4 galón',
        5, 5, date(2026, 10, 1), date(2026, 10, 1), None, 1,
    ])
    ns.append([
        2, date(2026, 10, 2), 'C-002', 'PIN-002', 'Innova — Celeste — balde 4 galones',
        9, 3, date(2026, 10, 2), date(2026, 10, 2), 'Sin stock', 0,
    ])
    # Fecha de atención como texto "dd/mm/aaaa" en vez de fecha nativa, para
    # probar que _a_datetime_aware tolera columnas mixtas del Excel real.
    ns.append([
        3, date(2026, 10, 3), 'C-003', None, 'Producto Nuevo — Rojo — galón',
        2, 2, date(2026, 10, 3), '03/10/2026', None, 1,
    ])

    # 3_COI_CA trae una columna "Periodo" entre "Detalle del cálculo" y
    # "Monto (S/)" en el archivo real; el comando no la lee, pero se incluye
    # para que la estructura coincida.
    coi_ca = libro.create_sheet('3_COI_CA')
    for _ in range(3):
        coi_ca.append([])
    coi_ca.append(['Componente', 'Detalle del cálculo', 'Periodo', 'Monto (S/)'])
    coi_ca.append(['Alquiler imputado al almacén', 'detalle x', '25/07/2026–24/08/2026', 100])
    coi_ca.append(['Mermas del periodo', 'detalle y', '25/07/2026–24/08/2026', 50])
    coi_ca.append(['Personal de almacén', 'detalle z', '25/07/2026–24/08/2026', 200])
    coi_ca.append(['TOTAL DE COSTOS DE ALMACENAMIENTO (CA)', None, None, 350])

    # 3_COI_PD trae "Código del producto" en el archivo real (columna
    # obligatoria para el comando: es la que usa para actualizar
    # precio_venta/costo_compra). La fila del pedido inexistente deja el
    # código en blanco a propósito, no afecta la advertencia que se prueba.
    coi_pd = libro.create_sheet('3_COI_PD')
    for _ in range(4):
        coi_pd.append([])
    coi_pd.append([
        'N.°', 'N.° de pedido en Ficha 2', 'Código del producto', 'Producto y presentación',
        'Cant. no atendida', 'Precio de venta (S/)', 'Costo de compra (S/)',
        'Margen unitario (S/)', 'Pérdida (S/)',
    ])
    coi_pd.append([1, 2, 'PIN-002', 'Innova — Celeste — balde 4 galones', 6, 10.0, 5.0, 5.0, 30.0])
    coi_pd.append([2, 999, None, 'Producto inexistente', 3, 1, 1, 0, 0])
    coi_pd.append([3, 1, 'PIN-001', 'CPP Pato — Blanco — 1/4 galón', 5, 1, 1, 0, 0])
    coi_pd.append(['TOTAL DE PÉRDIDAS POR DESABASTECIMIENTO (PD)', None, None, None, None, None, None, None, 30])

    libro.save(ruta)


class CargarDatosCommandTests(TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp_dir.cleanup)
        self.archivo = Path(self.tmp_dir.name) / 'ficha_prueba.xlsx'
        _construir_excel_minimo(self.archivo)

    def _ejecutar(self, **opciones):
        return call_command(
            'cargar_datos', archivo=str(self.archivo), origen='prueba',
            stdout=StringIO(), stderr=StringIO(), **opciones,
        )

    def test_dry_run_no_escribe_nada(self):
        self._ejecutar(dry_run=True)

        self.assertEqual(Producto.objects.count(), 0)
        self.assertEqual(Movimiento.objects.count(), 0)
        self.assertEqual(ConteoFisico.objects.count(), 0)
        self.assertEqual(ConteoDetalle.objects.count(), 0)
        self.assertEqual(Pedido.objects.count(), 0)
        self.assertEqual(PedidoDetalle.objects.count(), 0)
        self.assertEqual(CostoAlmacenamiento.objects.count(), 0)

    def test_importa_datos_correctos(self):
        salida = StringIO()
        call_command(
            'cargar_datos', archivo=str(self.archivo), origen='prueba',
            stdout=salida, stderr=StringIO(),
        )
        texto_salida = salida.getvalue()

        # 1_EI: PIN-001, PIN-002, PIN-003 + "Producto Nuevo" creado desde 2_NS.
        self.assertEqual(Producto.objects.count(), 4)
        self.assertEqual(Movimiento.objects.count(), 3)
        self.assertEqual(ConteoFisico.objects.count(), 1)
        self.assertEqual(ConteoDetalle.objects.count(), 3)
        self.assertEqual(Pedido.objects.count(), 3)
        self.assertEqual(PedidoDetalle.objects.count(), 3)
        # 3_COI_CA: 3 componentes, la fila TOTAL se omite.
        self.assertEqual(CostoAlmacenamiento.objects.count(), 3)

        # El ajuste de stock usa el valor del kardex, no el del conteo físico.
        pin001 = Producto.objects.get(codigo='PIN-001')
        self.assertEqual(pin001.stock_actual, 10)
        pin002 = Producto.objects.get(codigo='PIN-002')
        self.assertEqual(pin002.stock_actual, 9)

        # "Producto y presentación" separado en nombre/color/presentación.
        self.assertEqual(pin001.nombre, 'CPP Pato')
        self.assertEqual(pin001.color, 'Blanco')
        self.assertEqual(pin001.presentacion, '1/4 galón')

        # PIN-003 no tiene el formato "Marca — Color — Presentación".
        pin003 = Producto.objects.get(codigo='PIN-003')
        self.assertEqual(pin003.nombre, 'Producto sin formato válido')
        self.assertEqual(pin003.color, '')
        self.assertEqual(pin003.presentacion, '')

        # Producto nuevo detectado en 2_NS, sin código en el Excel.
        producto_nuevo = Producto.objects.get(nombre='Producto Nuevo')
        self.assertTrue(producto_nuevo.codigo.startswith('AUTO-'))

        # Pedido con atención parcial: motivo_no_atencion = sin_stock.
        detalle_parcial = PedidoDetalle.objects.get(cantidad_atendida=3)
        self.assertEqual(detalle_parcial.motivo_no_atencion, PedidoDetalle.MotivoNoAtencion.SIN_STOCK)
        self.assertFalse(detalle_parcial.atendido_a_tiempo)

        # "Fecha de atención" del pedido 3 vino como texto "03/10/2026": debe
        # interpretarse igual que si hubiera venido como fecha nativa.
        detalle_texto = PedidoDetalle.objects.get(producto__nombre='Producto Nuevo')
        self.assertEqual(detalle_texto.fecha_atencion.date(), date(2026, 10, 3))

        # Advertencias esperadas: producto sin formato, producto sin código,
        # pedido inexistente referenciado en 3_COI_PD y una discrepancia de cantidad.
        self.assertIn('no se pudo separar', texto_salida.lower())
        self.assertIn('no traía código', texto_salida)
        self.assertIn('N.° 999', texto_salida)
        self.assertIn('no coincide', texto_salida)

        # La fecha de atención en texto sí se pudo interpretar: no cuenta
        # como fecha inválida.
        self.assertIn('Fechas no interpretadas: 0', texto_salida)

        # Indicadores calculados desde la base de datos.
        self.assertIn('EI (exactitud del inventario): 66.67%', texto_salida)
        self.assertIn('NS (nivel de servicio):        66.67%', texto_salida)
        self.assertIn('S/ 350.00', texto_salida)

    def test_fecha_invalida_no_rompe_la_carga(self):
        # "Fecha de atención" del pedido 1 llega ilegible: la carga no debe
        # fallar, la fecha debe quedar vacía y debe reportarse en el resumen.
        libro = openpyxl.load_workbook(self.archivo)
        libro['2_NS'].cell(row=12, column=9).value = 'fecha-invalida'
        libro.save(self.archivo)

        salida = StringIO()
        call_command(
            'cargar_datos', archivo=str(self.archivo), origen='prueba',
            stdout=salida, stderr=StringIO(),
        )
        texto_salida = salida.getvalue()

        detalle = PedidoDetalle.objects.get(pedido__cliente='C-001')
        self.assertIsNone(detalle.fecha_atencion)
        self.assertIn('Fechas no interpretadas: 1', texto_salida)
        self.assertIn('2_NS fila 12, columna "Fecha de atención"', texto_salida)

    def test_limpiar_evita_duplicados(self):
        self._ejecutar()
        self.assertEqual(Producto.objects.count(), 4)
        self.assertEqual(Pedido.objects.count(), 3)

        self._ejecutar(limpiar=True)
        self.assertEqual(Producto.objects.count(), 4)
        self.assertEqual(Pedido.objects.count(), 3)
        self.assertEqual(Movimiento.objects.count(), 3)
