"""Reconstrucción del catálogo de productos desde la ficha de registro.

La hoja 1_EI de la ficha (el conteo físico del periodo) es la fuente
oficial del catálogo vigente: un producto existe si está en ese conteo.
`planificar()` compara la ficha con el catálogo y devuelve, sin escribir
nada, cada cambio campo por campo; `aplicar()` los escribe. Así la tabla de
cambios del anexo sale del mismo cálculo que se aplicó, no de una lista
hecha a mano.

Reglas (decididas por el usuario el 29/10/2026):
- nombre, color y presentación: los de la descripción de 1_EI;
- precio y costo: los de 3_COI_PD para los productos que aparecen ahí; los
  demás conservan los del catálogo;
- categoría: la del catálogo, emparejando por código; si es "otro" se
  deduce del nombre nuevo (clasificar_producto). `categorias_forzadas`
  permite fijar una a mano (p. ej. PIN-260, que pasó de Anypsa a Lija);
- productos del catálogo que no están en 1_EI: inactivos, nunca borrados.
"""
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal

import openpyxl

from ..management.commands.clasificar_productos import clasificar_producto
from ..models import Categoria, Producto

# Presentaciones reconocidas al final de la descripción, de la más larga a
# la más corta para que "balde 1 galón" no se lea como "galón".
PRESENTACIONES = (
    'balde 4 galones', 'balde 1 galón', '1/16 galón', '1/4 galón',
    'galón', 'litro', 'pliego', 'unidad', 'bolsa',
)

# Línea de producto (marca + línea) al inicio de la descripción, la más
# larga primero: "Anypsa Gloss" antes que "Anypsa", "Rodillo Toro" antes
# que "Rodillo". Son los nombres que ya usa el catálogo.
LINEAS = sorted((
    '3T Master Gloss', 'American Colors Mate', 'American Colors Satinado',
    'Anypsa', 'Anypsa Gloss', 'Base de rodillo', 'Bravazo Super Gloss',
    'Brocha C&A', 'Brocha Tumi', 'CPP Dura Látex', 'CPP Pato',
    'Ceresita Látex', 'Espátula', 'Galax Color', 'Gamax Satinado Color',
    'Gamax Satinado Paint', 'Innova', 'Lija', 'Rodillo', 'Rodillo Toro',
    'Temple Losaro', 'Temple Majestad', 'Velsalit', 'Vencedor Látex Premium',
), key=len, reverse=True)

CAMPOS_DESCRIPCION = ('nombre', 'color', 'presentacion')


def descripcion(producto):
    """'nombre color presentación', como lo escribe la ficha."""
    return ' '.join(filter(None, (producto.nombre, producto.color, producto.presentacion)))


def separar_descripcion(texto):
    """'Anypsa Gloss Azul honda balde 1 galón' -> ('Anypsa Gloss',
    'Azul honda', 'balde 1 galón'). Devuelve None si no empieza por una
    línea conocida o no termina en una presentación conocida: en ese caso no
    se adivina, se reporta."""
    texto = ' '.join(str(texto).split())
    presentacion = next((p for p in PRESENTACIONES if texto == p or texto.endswith(' ' + p)), None)
    linea = next((n for n in LINEAS if texto == n or texto.startswith(n + ' ')), None)
    if presentacion is None or linea is None:
        return None
    color = texto[len(linea):len(texto) - len(presentacion)].strip()
    if ' '.join(filter(None, (linea, color, presentacion))) != texto:
        return None
    return linea, color, presentacion


@dataclass
class Plan:
    """Lo que cambiaría la reconstrucción. cambios: [(codigo, campo,
    anterior, nuevo)] en el orden en que se aplican."""
    cambios: list = field(default_factory=list)
    inactivados: list = field(default_factory=list)       # [(codigo, descripción)]
    sin_separar: list = field(default_factory=list)       # [(codigo, descripción de la ficha)]
    faltan_en_catalogo: list = field(default_factory=list)
    categorias_deducidas: list = field(default_factory=list)   # [(codigo, descripción, categoría)]
    # Categoría conservada que no coincide con la que se deduce del nombre
    # nuevo: no se cambia (la regla es conservarla), solo se reporta.
    categorias_en_conflicto: list = field(default_factory=list)  # [(codigo, descripción, conservada, deducida)]
    duplicados: dict = field(default_factory=dict)        # {descripción: [códigos activos]}


def _leer_ficha(ruta):
    libro = openpyxl.load_workbook(ruta, data_only=True)
    filas_ei = [
        fila for fila in libro['1_EI'].iter_rows(values_only=True)
        if isinstance(fila[0], int) and fila[2]
    ]
    descripciones = {str(fila[2]).strip(): ' '.join(str(fila[3]).split()) for fila in filas_ei}
    precios = {}
    for fila in libro['3_COI_PD'].iter_rows(values_only=True):
        if isinstance(fila[0], int) and fila[2]:
            precios[str(fila[2]).strip()] = (Decimal(str(fila[5])), Decimal(str(fila[6])))
    return descripciones, precios


def planificar(ruta, categorias_forzadas=None):
    """Compara la ficha con el catálogo actual. No escribe nada."""
    categorias_forzadas = categorias_forzadas or {}
    descripciones, precios = _leer_ficha(ruta)
    catalogo = {p.codigo: p for p in Producto.objects.all()}
    plan = Plan()

    nuevos = {}
    for codigo, texto in descripciones.items():
        producto = catalogo.get(codigo)
        if producto is None:
            plan.faltan_en_catalogo.append((codigo, texto))
            continue
        partes = separar_descripcion(texto)
        if partes is None:
            plan.sin_separar.append((codigo, texto))
            continue
        valores = dict(zip(CAMPOS_DESCRIPCION, partes))
        if codigo in precios:
            valores['precio_venta'], valores['costo_compra'] = precios[codigo]

        # La categoría se evalúa con el nombre nuevo, no con el del catálogo.
        provisional = Producto(**{c: valores[c] for c in CAMPOS_DESCRIPCION})
        deducida = clasificar_producto(provisional)
        if codigo in categorias_forzadas:
            valores['categoria'] = categorias_forzadas[codigo]
        elif producto.categoria == Categoria.OTRO:
            valores['categoria'] = deducida
            plan.categorias_deducidas.append((codigo, texto, deducida))
        elif deducida not in (producto.categoria, Categoria.OTRO):
            plan.categorias_en_conflicto.append((codigo, texto, producto.categoria, deducida))
        if not producto.activo:
            valores['activo'] = True
        nuevos[codigo] = valores

        for campo, valor in valores.items():
            anterior = getattr(producto, campo)
            if anterior != valor:
                plan.cambios.append((codigo, campo, anterior, valor))

    for codigo, producto in sorted(catalogo.items()):
        if codigo not in descripciones and producto.activo:
            plan.inactivados.append((codigo, descripcion(producto)))
            plan.cambios.append((codigo, 'activo', True, False))

    grupos = defaultdict(list)
    for codigo, texto in descripciones.items():
        if codigo in nuevos:
            grupos[texto].append(codigo)
    plan.duplicados = {texto: sorted(c) for texto, c in sorted(grupos.items()) if len(c) > 1}
    return plan


def aplicar(plan):
    """Escribe los cambios del plan en el catálogo. Debe llamarse dentro de
    una transacción (el comando la abre)."""
    por_codigo = defaultdict(dict)
    for codigo, campo, _anterior, nuevo in plan.cambios:
        por_codigo[codigo][campo] = nuevo
    for codigo, valores in por_codigo.items():
        producto = Producto.objects.get(codigo=codigo)
        for campo, valor in valores.items():
            setattr(producto, campo, valor)
        producto.save(update_fields=[*valores, 'actualizado_en'])
    return len(por_codigo)
