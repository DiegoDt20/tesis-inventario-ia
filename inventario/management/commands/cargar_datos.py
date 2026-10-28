"""Importa datos desde el Excel de fichas de registro (EI, NS, COI) al sistema."""
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import from_excel
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from inventario.servicios.indicadores import calcular_coi, calcular_ei, calcular_ns
from inventario.models import (
    ConteoDetalle,
    ConteoFisico,
    CostoAlmacenamiento,
    Movimiento,
    Origen,
    Pedido,
    PedidoDetalle,
    Producto,
)

# Guion largo usado en el Excel para separar "Marca — Color — Presentación".
SEPARADOR_PRODUCTO = '—'

USUARIO_IMPORTACION = 'importador_datos'

# Formatos de texto aceptados para fechas que llegaron como cadena en vez de
# como celda de fecha nativa de Excel.
FORMATOS_FECHA_TEXTO = ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y')

# Periodo declarado en la columna "Periodo" de 3_COI_CA, p. ej.
# "25/07/2026–24/08/2026" (guion largo, corto o "al" entre las dos fechas).
PATRON_PERIODO = re.compile(
    r'(\d{1,2}/\d{1,2}/\d{4})\s*(?:[–—-]|al)\s*(\d{1,2}/\d{1,2}/\d{4})'
)


class FilaInvalida(Exception):
    """Señala que una fila debe omitirse porque algún valor no tiene el tipo
    esperado (p. ej. texto donde se esperaba un número). El detalle ya quedó
    registrado como advertencia antes de lanzar esta excepción."""


@dataclass
class Resumen:
    """Acumula lo que hizo la importación para el reporte final."""
    productos_creados: int = 0
    pedidos_creados: int = 0
    movimientos_creados: int = 0
    conteos_creados: int = 0
    costos_creados: int = 0
    filas_omitidas: int = 0
    advertencias: list = field(default_factory=list)
    fechas_invalidas: list = field(default_factory=list)
    # Meses (primer día) en los que este archivo dejó datos de algún
    # indicador: conteos (EI), pedidos (NS y PD) y costos (CA). Sirven para
    # no mezclar periodos (ver Command._verificar_periodos) y para reportar
    # los indicadores solo del periodo cargado.
    meses_cargados: set = field(default_factory=set)

    def advertir(self, texto):
        self.advertencias.append(texto)

    def registrar_fecha_invalida(self, hoja, fila, columna, valor):
        self.fechas_invalidas.append((hoja, fila, columna, valor))
        self.advertir(
            f'{hoja} fila {fila}, columna "{columna}": no se pudo interpretar '
            f'la fecha "{valor}", se dejó vacía.'
        )


class Command(BaseCommand):
    help = (
        'Importa los datos de las fichas de registro (Excel) a la base de datos: '
        'exactitud de inventario (1_EI), nivel de servicio (2_NS) y costos '
        'operativos de inventario (3_COI_CA y 3_COI_PD).'
    )

    HOJAS_REQUERIDAS = ('1_EI', '2_NS', '3_COI_CA', '3_COI_PD')
    # Quinta hoja: solo resume lo ya calculado en las otras cuatro, con
    # cualquiera de estos dos nombres según la versión de la plantilla. No
    # aporta datos nuevos, así que si falta simplemente se omite.
    HOJAS_RESUMEN_COI = ('3_COI_registro', '3_COI_resumen')

    # Máximo de filas iniciales donde se busca la fila de encabezado de cada
    # hoja (las filas anteriores son título, fórmula, periodo, etc.).
    MAX_FILAS_BUSQUEDA_ENCABEZADO = 20

    def add_arguments(self, parser):
        parser.add_argument(
            '--archivo', required=True,
            help='Ruta al archivo Excel (.xlsx) con las fichas de registro.',
        )
        parser.add_argument(
            '--origen', required=True, choices=Origen.values,
            help='Origen de los datos a cargar: prueba o real.',
        )
        parser.add_argument(
            '--limpiar', action='store_true',
            help='Borra los datos previos del mismo origen antes de cargar.',
        )
        parser.add_argument(
            '--dry-run', action='store_true',
            help='Procesa y reporta el resultado sin escribir cambios en la base de datos.',
        )

    def handle(self, *args, **options):
        ruta = Path(options['archivo'])
        origen = options['origen']
        limpiar = options['limpiar']
        dry_run = options['dry_run']

        if not ruta.exists():
            raise CommandError(f'No se encontró el archivo: {ruta}')

        try:
            libro = openpyxl.load_workbook(ruta, data_only=True)
        except Exception as exc:
            raise CommandError(f'No se pudo abrir el archivo Excel: {exc}')

        faltantes = [h for h in self.HOJAS_REQUERIDAS if h not in libro.sheetnames]
        if faltantes:
            raise CommandError(
                f'Al archivo le faltan las hojas: {", ".join(faltantes)}'
            )

        resumen = Resumen()

        if not any(nombre in libro.sheetnames for nombre in self.HOJAS_RESUMEN_COI):
            resumen.advertir(
                'No se encontró ninguna hoja de resumen de COI ({}); se omite: '
                'es una hoja de resumen y no aporta datos nuevos.'.format(
                    ' ni '.join(f'"{n}"' for n in self.HOJAS_RESUMEN_COI)
                )
            )

        # Meses que ya tenían datos de este origen ANTES de la carga: con
        # ellos se impide que el archivo se mezcle con un periodo ya medido
        # (p. ej. el postest con el pretest), ver _verificar_periodos.
        meses_previos = self._meses_con_datos(origen)

        with transaction.atomic():
            if limpiar:
                self._limpiar_origen(origen)

            usuario = self._obtener_usuario_importacion()
            productos_por_codigo = {}
            productos_por_descripcion = {}

            # 3_COI_PD se lee antes que 2_NS porque trae el precio y el costo
            # de cada línea no atendida en el periodo de la ficha: son los
            # que se congelan en PedidoDetalle al crearla (no el precio que
            # el producto tenga en el catálogo al momento de importar).
            filas_pd = self._leer_coi_pd(libro['3_COI_PD'], resumen)
            precios_por_pedido = {}
            for fila_pd in filas_pd:
                precios_por_pedido.setdefault(
                    fila_pd['numero_pedido'], (fila_pd['precio_venta'], fila_pd['costo_compra']),
                )

            fecha_corte_ei = self._cargar_ei(
                libro['1_EI'], origen, usuario, productos_por_codigo,
                productos_por_descripcion, resumen,
            )
            pedidos_por_numero = self._cargar_ns(
                libro['2_NS'], origen, productos_por_codigo,
                productos_por_descripcion, resumen, precios_por_pedido,
            )
            self._cargar_coi_ca(
                libro['3_COI_CA'], origen, fecha_corte_ei, resumen, pedidos_por_numero,
            )
            self._cargar_coi_pd(filas_pd, pedidos_por_numero, productos_por_codigo, resumen)

            self._verificar_periodos(origen, meses_previos, resumen.meses_cargados, limpiar)

            self._imprimir_resumen(resumen)
            self._imprimir_indicadores(origen, resumen.meses_cargados)

            if dry_run:
                transaction.set_rollback(True)
                self.stdout.write(self.style.WARNING(
                    '\nDry-run: no se escribió ningún cambio en la base de datos.'
                ))

    # ------------------------------------------------------------------
    # Lectura robusta de hojas: encabezados por texto, no por posición
    # ------------------------------------------------------------------

    @staticmethod
    def _normalizar_texto(valor):
        """minúsculas, sin tildes, sin signos de puntuación, sin espacios
        dobles. Sirve para comparar encabezados sin depender de mayúsculas,
        acentos o puntuación (p. ej. "N.°" y "Código" -> "n" y "codigo")."""
        if valor is None:
            return ''
        texto = str(valor).strip().lower()
        texto = unicodedata.normalize('NFKD', texto)
        texto = ''.join(c for c in texto if not unicodedata.combining(c))
        texto = re.sub(r'[^a-z0-9\s]', ' ', texto)
        texto = re.sub(r'\s+', ' ', texto).strip()
        return texto

    def _encontrar_fila_encabezado(self, hoja, nombre_hoja, marcador='n'):
        """Busca en las primeras filas de la hoja aquella cuya primera
        columna, normalizada, sea igual a `marcador` (por defecto "n", que es
        cómo queda "N.°" tras normalizar). Los datos empiezan en la fila
        siguiente. No asume una posición fija porque cambia entre fichas."""
        limite = min(self.MAX_FILAS_BUSQUEDA_ENCABEZADO, hoja.max_row)
        for fila in range(1, limite + 1):
            texto = self._normalizar_texto(hoja.cell(row=fila, column=1).value)
            if texto == marcador:
                return fila
        raise CommandError(
            f'{nombre_hoja}: no se encontró la fila de encabezado (se buscó '
            f'"{marcador}" en la primera columna de las primeras {limite} filas).'
        )

    def _mapear_columnas(self, hoja, fila_encabezado, alias_por_columna, requeridas, nombre_hoja):
        """Lee la fila de encabezado y arma un diccionario
        nombre_de_columna -> índice de columna, comparando el texto
        normalizado de cada celda contra los alias conocidos de cada campo.
        Falla con un mensaje claro si falta una columna obligatoria."""
        encontrados = {}
        for columna in range(1, hoja.max_column + 1):
            texto = self._normalizar_texto(hoja.cell(row=fila_encabezado, column=columna).value)
            if not texto:
                continue
            for nombre_canonico, alias in alias_por_columna.items():
                if nombre_canonico not in encontrados and texto in alias:
                    encontrados[nombre_canonico] = columna

        faltantes = [nombre for nombre in requeridas if nombre not in encontrados]
        if faltantes:
            raise CommandError(
                f'{nombre_hoja}: faltan columnas obligatorias {faltantes} en la fila '
                f'de encabezado (fila {fila_encabezado}). Columnas encontradas: '
                f'{sorted(encontrados)}.'
            )
        return encontrados

    @staticmethod
    def _es_numerico(valor):
        if isinstance(valor, bool):
            return True
        return isinstance(valor, (int, float))

    def _a_entero_columna(self, resumen, hoja_nombre, fila, columna_nombre, valor):
        """Convierte a int validando el tipo antes; si el valor no es
        numérico (p. ej. vino como texto), registra la advertencia con hoja,
        fila, columna y valor, y señala que la fila debe omitirse."""
        if not self._es_numerico(valor):
            resumen.advertir(
                f'{hoja_nombre} fila {fila}, columna "{columna_nombre}": se esperaba '
                f'un valor numérico y se encontró "{valor}", se omitió la fila.'
            )
            raise FilaInvalida
        return int(valor)

    def _a_decimal_columna(self, resumen, hoja_nombre, fila, columna_nombre, valor):
        if not self._es_numerico(valor):
            resumen.advertir(
                f'{hoja_nombre} fila {fila}, columna "{columna_nombre}": se esperaba '
                f'un valor numérico y se encontró "{valor}", se omitió la fila.'
            )
            raise FilaInvalida
        return Decimal(str(valor))

    # ------------------------------------------------------------------
    # Utilidades generales
    # ------------------------------------------------------------------

    def _obtener_usuario_importacion(self):
        """Usuario técnico dueño de los movimientos de ajuste generados por
        la importación (Movimiento.usuario es obligatorio)."""
        Usuario = get_user_model()
        usuario, creado = Usuario.objects.get_or_create(
            username=USUARIO_IMPORTACION,
            defaults={'first_name': 'Importación de datos', 'is_active': False},
        )
        if creado:
            usuario.set_unusable_password()
            usuario.save(update_fields=['password'])
        return usuario

    @staticmethod
    def _meses_con_datos(origen):
        """Meses (primer día) con datos de algún indicador del origen:
        conteos físicos (EI), pedidos (NS y PD) y costos de almacenamiento
        (CA). Los movimientos no cuentan: no entran en ningún indicador."""
        meses = set(
            ConteoFisico.objects.filter(origen=origen).dates('fecha_corte', 'month')
        )
        meses.update(
            momento.date() for momento in Pedido.objects.filter(origen=origen).datetimes(
                'fecha_solicitud', 'month', tzinfo=timezone.get_current_timezone(),
            )
        )
        meses.update(
            CostoAlmacenamiento.objects.filter(origen=origen).dates('periodo_mes', 'month')
        )
        return meses

    @staticmethod
    def _verificar_periodos(origen, meses_previos, meses_cargados, limpiar):
        """Impide que una carga cambie un periodo ya medido de este origen
        (lo que congela el pretest al importar el postest). Lanza
        CommandError dentro de la transacción, así que no queda nada
        escrito:

        - sin --limpiar, el archivo no puede traer datos de un mes que ya
          tenía datos del mismo origen (se sumarían a ese periodo y
          cambiarían sus indicadores);
        - con --limpiar, que borra TODO el origen, el archivo tiene que
          volver a cargar todos los meses que había; si no, --limpiar
          estaría borrando otro periodo (p. ej. el pretest al cargar el
          postest con --limpiar).
        """
        def formatear(meses):
            return ', '.join(m.strftime('%m/%Y') for m in sorted(meses))

        if limpiar:
            perdidos = meses_previos - meses_cargados
            if perdidos:
                raise CommandError(
                    f'--limpiar borraría los datos de origen "{origen}" de {formatear(perdidos)}, '
                    f'que este archivo no vuelve a cargar (solo trae {formatear(meses_cargados) or "nada"}). '
                    'No se escribió nada. Para cargar un periodo nuevo sin tocar los '
                    'anteriores, ejecútalo sin --limpiar.'
                )
            return

        solapados = meses_previos & meses_cargados
        if solapados:
            raise CommandError(
                f'El archivo trae datos de {formatear(solapados)}, y ya hay datos de origen '
                f'"{origen}" en ese mes: se sumarían a ese periodo y cambiarían sus '
                'indicadores. No se escribió nada. Revisa las fechas del archivo (p. ej. '
                'pedidos del postest con fecha del pretest) o, si de verdad quieres '
                'reemplazar ese periodo, restaura el respaldo previo a su carga.'
            )

    def _limpiar_origen(self, origen):
        """Borra los datos previos del origen indicado, respetando el orden
        de dependencias (los detalles antes que sus cabeceras/productos)."""
        Movimiento.objects.filter(origen=origen).delete()
        PedidoDetalle.objects.filter(pedido__origen=origen).delete()
        Pedido.objects.filter(origen=origen).delete()
        ConteoDetalle.objects.filter(conteo__origen=origen).delete()
        ConteoFisico.objects.filter(origen=origen).delete()
        CostoAlmacenamiento.objects.filter(origen=origen).delete()
        Producto.objects.filter(origen=origen).delete()

    @staticmethod
    def _parsear_valor_fecha(valor):
        """Intenta interpretar el valor de una celda como fecha/hora, sin
        importar si Excel la entregó como datetime, texto o número serial.
        Devuelve un datetime (naive) o None si no se pudo interpretar."""
        if valor is None:
            return None
        if isinstance(valor, datetime):
            return valor
        if isinstance(valor, date):
            return datetime.combine(valor, time.min)
        if isinstance(valor, str):
            texto = valor.strip()
            if not texto:
                return None
            for formato in FORMATOS_FECHA_TEXTO:
                try:
                    return datetime.strptime(texto, formato)
                except ValueError:
                    continue
            return None
        if isinstance(valor, (int, float)):
            try:
                return from_excel(valor)
            except (ValueError, TypeError):
                return None
        return None

    def _a_fecha(self, valor, resumen, hoja, fila, columna):
        """Normaliza una celda de fecha del Excel a date, tolerando texto,
        datetime o serial de Excel. Registra advertencia si no vino vacía
        pero tampoco se pudo interpretar."""
        resultado = self._parsear_valor_fecha(valor)
        if resultado is not None:
            return resultado.date()
        if valor is not None and str(valor).strip():
            resumen.registrar_fecha_invalida(hoja, fila, columna, valor)
        return None

    def _a_datetime_aware(self, valor, resumen, hoja, fila, columna):
        """Normaliza una celda de fecha/hora del Excel a datetime con zona
        horaria, con la misma tolerancia que _a_fecha."""
        resultado = self._parsear_valor_fecha(valor)
        if resultado is None:
            if valor is not None and str(valor).strip():
                resumen.registrar_fecha_invalida(hoja, fila, columna, valor)
            return None
        if timezone.is_naive(resultado):
            resultado = timezone.make_aware(resultado)
        return resultado

    @staticmethod
    def _fecha_a_datetime_aware(fecha):
        """Convierte una date ya validada (p. ej. fecha_corte) a datetime con
        zona horaria, sin volver a intentar parsearla."""
        if fecha is None:
            return None
        resultado = datetime.combine(fecha, time.min)
        return timezone.make_aware(resultado) if timezone.is_naive(resultado) else resultado

    def _parsear_producto_presentacion(self, texto):
        """Separa "Marca — Color — Presentación" en (nombre, color,
        presentacion, ok). Si no hay exactamente 3 partes, todo va a nombre."""
        texto = (texto or '').strip()
        partes = [p.strip() for p in texto.split(SEPARADOR_PRODUCTO)]
        if len(partes) == 3 and all(partes):
            return partes[0], partes[1], partes[2], True
        return texto, '', '', False

    def _obtener_o_crear_producto(
        self, codigo, descripcion, origen, productos_por_codigo,
        productos_por_descripcion, resumen,
    ):
        """Busca un producto por código (o por descripción si no hay código)
        y lo crea si no existe, usando los datos disponibles en la fila."""
        descripcion_normalizada = (descripcion or '').strip()

        if codigo:
            codigo = str(codigo).strip()
            producto = productos_por_codigo.get(codigo)
            if producto is None:
                producto = Producto.objects.filter(codigo=codigo).first()
            if producto:
                productos_por_codigo[codigo] = producto
                productos_por_descripcion.setdefault(descripcion_normalizada, producto)
                return producto
        else:
            producto = productos_por_descripcion.get(descripcion_normalizada)
            if producto:
                return producto

        nombre, color, presentacion, ok = self._parsear_producto_presentacion(descripcion)
        if not ok:
            resumen.advertir(
                f'No se pudo separar "{descripcion}" en marca, color y presentación '
                f'(se guardó todo en el campo nombre).'
            )

        if not codigo:
            codigo = f'AUTO-{Producto.objects.count() + 1:04d}'
            resumen.advertir(
                f'El producto "{descripcion}" no traía código (hoja 2_NS); '
                f'se generó el código "{codigo}" automáticamente.'
            )

        producto = Producto.objects.create(
            codigo=codigo,
            nombre=nombre,
            color=color,
            presentacion=presentacion,
            precio_venta=Decimal('0.00'),
            costo_compra=Decimal('0.00'),
            origen=origen,
        )
        resumen.productos_creados += 1
        productos_por_codigo[codigo] = producto
        productos_por_descripcion.setdefault(descripcion_normalizada, producto)
        return producto

    # ------------------------------------------------------------------
    # 1_EI - Exactitud del inventario
    # ------------------------------------------------------------------

    def _cargar_ei(
        self, hoja, origen, usuario, productos_por_codigo,
        productos_por_descripcion, resumen,
    ):
        nombre_hoja = '1_EI'
        fila_encabezado = self._encontrar_fila_encabezado(hoja, nombre_hoja)
        columnas = self._mapear_columnas(
            hoja, fila_encabezado,
            alias_por_columna={
                'numero': ['n'],
                'fecha_corte': ['fecha de corte'],
                'codigo': ['codigo'],
                'descripcion': ['producto y presentacion'],
                'stock_kardex': ['stock segun kardex'],
                'stock_fisico': ['stock segun conteo fisico'],
            },
            requeridas=[
                'numero', 'fecha_corte', 'codigo', 'descripcion',
                'stock_kardex', 'stock_fisico',
            ],
            nombre_hoja=nombre_hoja,
        )

        conteos_por_fecha = {}
        ultima_fecha_corte = None

        for fila in range(fila_encabezado + 1, hoja.max_row + 1):
            numero = hoja.cell(row=fila, column=columnas['numero']).value
            if numero is None:
                break

            fecha_corte = self._a_fecha(
                hoja.cell(row=fila, column=columnas['fecha_corte']).value,
                resumen, nombre_hoja, fila, 'Fecha de corte',
            )
            codigo = hoja.cell(row=fila, column=columnas['codigo']).value
            descripcion = hoja.cell(row=fila, column=columnas['descripcion']).value

            if fecha_corte is None or codigo is None:
                resumen.filas_omitidas += 1
                resumen.advertir(f'{nombre_hoja} fila {fila}: falta fecha de corte o código, se omitió.')
                continue

            try:
                stock_kardex = self._a_entero_columna(
                    resumen, nombre_hoja, fila, 'Stock según kardex',
                    hoja.cell(row=fila, column=columnas['stock_kardex']).value,
                )
                stock_fisico = self._a_entero_columna(
                    resumen, nombre_hoja, fila, 'Stock según conteo físico',
                    hoja.cell(row=fila, column=columnas['stock_fisico']).value,
                )
            except FilaInvalida:
                resumen.filas_omitidas += 1
                continue

            producto = self._obtener_o_crear_producto(
                codigo, descripcion, origen, productos_por_codigo,
                productos_por_descripcion, resumen,
            )

            conteo = conteos_por_fecha.get(fecha_corte)
            if conteo is None:
                conteo = ConteoFisico.objects.create(
                    fecha_corte=fecha_corte,
                    responsable='Importación automática',
                    origen=origen,
                )
                conteos_por_fecha[fecha_corte] = conteo
                resumen.conteos_creados += 1
                resumen.meses_cargados.add(fecha_corte.replace(day=1))

            ConteoDetalle.objects.create(
                conteo=conteo,
                producto=producto,
                stock_sistema=stock_kardex,
                stock_fisico=stock_fisico,
            )

            # No se escribe stock_actual a mano: se crea un movimiento de
            # ajuste que fija el stock al valor del kardex y deja que la
            # lógica de Movimiento.save() actualice Producto.stock_actual.
            Movimiento.objects.create(
                producto=producto,
                tipo=Movimiento.Tipo.AJUSTE,
                cantidad=stock_kardex,
                fecha=self._fecha_a_datetime_aware(fecha_corte),
                motivo=f'Ajuste por conteo físico importado (1_EI fila {fila})',
                usuario=usuario,
                origen=origen,
            )
            resumen.movimientos_creados += 1
            ultima_fecha_corte = fecha_corte

        return ultima_fecha_corte

    # ------------------------------------------------------------------
    # 2_NS - Nivel de servicio
    # ------------------------------------------------------------------

    def _cargar_ns(
        self, hoja, origen, productos_por_codigo, productos_por_descripcion, resumen,
        precios_por_pedido=None,
    ):
        """Crea un Pedido con su PedidoDetalle por fila. precios_por_pedido
        ({N.° de pedido: (precio_venta, costo_compra)}, leído de 3_COI_PD)
        fija el precio y el costo congelados de las líneas no atendidas por
        completo; las demás congelan el precio del producto al importar
        (no entran en el PD)."""
        precios_por_pedido = precios_por_pedido or {}
        nombre_hoja = '2_NS'
        fila_encabezado = self._encontrar_fila_encabezado(hoja, nombre_hoja)
        columnas = self._mapear_columnas(
            hoja, fila_encabezado,
            alias_por_columna={
                'numero': ['n'],
                'fecha_pedido': ['fecha del pedido'],
                'cliente': ['cliente codigo', 'cliente'],
                'codigo': ['codigo del producto'],
                'descripcion': ['producto y presentacion'],
                'cantidad_solicitada': ['cant solicitada', 'cantidad solicitada'],
                'cantidad_atendida': ['cant atendida', 'cantidad atendida'],
                'fecha_requerida': ['fecha requerida'],
                'fecha_atencion': ['fecha de atencion'],
                'a_tiempo': ['a tiempo 1 0', 'a tiempo'],
            },
            requeridas=[
                'numero', 'fecha_pedido', 'cliente', 'descripcion',
                'cantidad_solicitada', 'cantidad_atendida', 'fecha_atencion', 'a_tiempo',
            ],
            nombre_hoja=nombre_hoja,
        )

        pedidos_por_numero = {}

        for fila in range(fila_encabezado + 1, hoja.max_row + 1):
            numero = hoja.cell(row=fila, column=columnas['numero']).value
            if numero is None:
                break

            fecha_pedido = self._a_datetime_aware(
                hoja.cell(row=fila, column=columnas['fecha_pedido']).value,
                resumen, nombre_hoja, fila, 'Fecha del pedido',
            )
            cliente = hoja.cell(row=fila, column=columnas['cliente']).value
            descripcion = hoja.cell(row=fila, column=columnas['descripcion']).value
            # A diferencia de plantillas anteriores, esta hoja sí trae el
            # código del producto; se usa para emparejar con lo ya cargado
            # en 1_EI en vez de depender solo de la descripción.
            codigo = hoja.cell(row=fila, column=columnas['codigo']).value if 'codigo' in columnas else None
            cantidad_solicitada_valor = hoja.cell(row=fila, column=columnas['cantidad_solicitada']).value
            cantidad_atendida_valor = hoja.cell(row=fila, column=columnas['cantidad_atendida']).value
            fecha_atencion = self._a_datetime_aware(
                hoja.cell(row=fila, column=columnas['fecha_atencion']).value,
                resumen, nombre_hoja, fila, 'Fecha de atención',
            )
            a_tiempo = hoja.cell(row=fila, column=columnas['a_tiempo']).value
            fecha_requerida = (
                self._a_fecha(
                    hoja.cell(row=fila, column=columnas['fecha_requerida']).value,
                    resumen, nombre_hoja, fila, 'Fecha requerida',
                ) if 'fecha_requerida' in columnas else None
            )

            if fecha_pedido is None or cantidad_solicitada_valor is None:
                resumen.filas_omitidas += 1
                resumen.advertir(f'{nombre_hoja} fila {fila}: faltan datos obligatorios, se omitió.')
                continue

            try:
                cantidad_solicitada = self._a_entero_columna(
                    resumen, nombre_hoja, fila, 'Cant. solicitada', cantidad_solicitada_valor,
                )
                cantidad_atendida = (
                    self._a_entero_columna(
                        resumen, nombre_hoja, fila, 'Cant. atendida', cantidad_atendida_valor,
                    ) if cantidad_atendida_valor is not None else 0
                )
            except FilaInvalida:
                resumen.filas_omitidas += 1
                continue

            producto = self._obtener_o_crear_producto(
                codigo, descripcion, origen, productos_por_codigo,
                productos_por_descripcion, resumen,
            )

            motivo_no_atencion = None
            if cantidad_atendida < cantidad_solicitada:
                motivo_no_atencion = PedidoDetalle.MotivoNoAtencion.SIN_STOCK

            a_tiempo = self._resolver_a_tiempo(
                a_tiempo, cantidad_solicitada, cantidad_atendida, fecha_requerida,
                fecha_atencion, resumen, nombre_hoja, fila,
            )
            precio_venta_unitario, costo_compra_unitario = precios_por_pedido.get(numero, (None, None))

            pedido = Pedido.objects.create(
                fecha_solicitud=fecha_pedido,
                cliente=str(cliente or ''),
                # El Excel no registra el canal de venta; se asume "mostrador"
                # por ser el canal más habitual del negocio.
                canal=Pedido.Canal.MOSTRADOR,
                estado=(
                    Pedido.Estado.ATENDIDO if cantidad_atendida >= cantidad_solicitada
                    else Pedido.Estado.PENDIENTE
                ),
                origen=origen,
            )
            PedidoDetalle.objects.create(
                pedido=pedido,
                producto=producto,
                cantidad_solicitada=cantidad_solicitada,
                cantidad_atendida=cantidad_atendida,
                fecha_requerida=fecha_requerida,
                fecha_atencion=fecha_atencion,
                atendido_a_tiempo=a_tiempo,
                motivo_no_atencion=motivo_no_atencion,
                # None = el save() congela el precio actual del producto.
                precio_venta_unitario=precio_venta_unitario,
                costo_compra_unitario=costo_compra_unitario,
            )
            resumen.pedidos_creados += 1
            resumen.meses_cargados.add(timezone.localdate(fecha_pedido).replace(day=1))
            pedidos_por_numero[numero] = pedido

        return pedidos_por_numero

    @staticmethod
    def _resolver_a_tiempo(
        a_tiempo, solicitada, atendida, fecha_requerida, fecha_atencion,
        resumen, nombre_hoja, fila,
    ):
        """Valor de atendido_a_tiempo. Manda la columna "¿A tiempo? (1/0)"
        de la ficha; si la fecha requerida está, se contrasta con el
        criterio de la ficha (1 solo si se entrega la cantidad completa en
        la fecha requerida o antes) y una discrepancia queda como
        advertencia. Si la celda viene vacía (fórmula sin valor calculado,
        p. ej. un archivo generado por script que nunca se abrió en Excel),
        se usa el criterio en vez de contar el pedido como "no a tiempo"."""
        criterio = None
        if fecha_requerida is not None:
            criterio = (
                atendida >= solicitada
                and fecha_atencion is not None
                and timezone.localdate(fecha_atencion) <= fecha_requerida
            )
        if a_tiempo is None or a_tiempo == '':
            if criterio is None:
                return False
            resumen.advertir(
                f'{nombre_hoja} fila {fila}: "¿A tiempo?" vacío; se calculó con la '
                f'fecha requerida y la de atención ({int(criterio)}).'
            )
            return criterio
        valor = bool(a_tiempo)
        if criterio is not None and valor != criterio:
            resumen.advertir(
                f'{nombre_hoja} fila {fila}: "¿A tiempo?" dice {int(valor)}, pero con la '
                f'fecha requerida y la de atención el criterio da {int(criterio)}; se '
                'conservó el valor de la ficha.'
            )
        return valor

    # ------------------------------------------------------------------
    # 3_COI_CA - Costos de almacenamiento
    # ------------------------------------------------------------------

    def _cargar_coi_ca(self, hoja, origen, fecha_corte_ei, resumen, pedidos_por_numero=None):
        nombre_hoja = '3_COI_CA'
        # Esta hoja no numera filas con "N.°": su primera columna es
        # "Componente", así que se usa ese texto como marcador de encabezado.
        fila_encabezado = self._encontrar_fila_encabezado(hoja, nombre_hoja, marcador='componente')
        columnas = self._mapear_columnas(
            hoja, fila_encabezado,
            alias_por_columna={
                'componente': ['componente'],
                'detalle': ['detalle del calculo'],
                'periodo': ['periodo'],
                'monto': ['monto s', 'monto'],
            },
            requeridas=['componente', 'monto'],
            nombre_hoja=nombre_hoja,
        )

        # CostoAlmacenamiento solo guarda el mes (periodo_mes): se usa el
        # mes de la fecha de corte de 1_EI como referencia del periodo
        # evaluado. La ventana exacta que declara la columna "Periodo" no se
        # guarda; se contrasta con la de los pedidos (_contrastar_ventana_ca).
        if fecha_corte_ei is not None:
            periodo_mes = fecha_corte_ei.replace(day=1)
        else:
            periodo_mes = timezone.localdate().replace(day=1)
            resumen.advertir(
                f'{nombre_hoja}: no se pudo determinar el periodo desde 1_EI, '
                'se usó el mes actual para periodo_mes.'
            )

        periodos_declarados = set()
        for fila in range(fila_encabezado + 1, hoja.max_row + 1):
            componente = hoja.cell(row=fila, column=columnas['componente']).value
            if componente is not None and 'periodo' in columnas:
                periodo_texto = hoja.cell(row=fila, column=columnas['periodo']).value
                if periodo_texto:
                    periodos_declarados.add(str(periodo_texto).strip())
            detalle = (
                hoja.cell(row=fila, column=columnas['detalle']).value
                if 'detalle' in columnas else None
            )
            monto_valor = hoja.cell(row=fila, column=columnas['monto']).value

            if componente is None:
                break
            if str(componente).strip().upper().startswith('TOTAL'):
                continue
            if monto_valor is None:
                resumen.filas_omitidas += 1
                resumen.advertir(f'{nombre_hoja} fila {fila}: sin monto, se omitió.')
                continue

            try:
                monto = self._a_decimal_columna(resumen, nombre_hoja, fila, 'Monto (S/)', monto_valor)
            except FilaInvalida:
                resumen.filas_omitidas += 1
                continue

            concepto = str(componente).strip()
            if detalle:
                concepto = f'{concepto} — {str(detalle).strip()}'[:200]

            CostoAlmacenamiento.objects.create(
                periodo_mes=periodo_mes,
                concepto=concepto,
                monto=monto,
                origen=origen,
            )
            resumen.costos_creados += 1
            resumen.meses_cargados.add(periodo_mes)

            # Nota: "Mermas del periodo" solo se registra como costo de
            # almacenamiento porque el Excel no trae el desglose por
            # producto/cantidad que exige el modelo Merma.
            if str(componente).strip().lower().startswith('merma'):
                resumen.advertir(
                    f'{nombre_hoja} fila {fila}: "{concepto}" es un monto agregado '
                    f'sin desglose por producto, no se creó un registro Merma '
                    f'(solo el CostoAlmacenamiento correspondiente).'
                )

        self._contrastar_ventana_ca(periodos_declarados, pedidos_por_numero, periodo_mes, resumen)

    @staticmethod
    def _contrastar_ventana_ca(periodos_declarados, pedidos_por_numero, periodo_mes, resumen):
        """El COI suma CA + PD, así que ambos deben medir la misma ventana.
        CostoAlmacenamiento solo guarda el mes (periodo_mes), así que la
        ventana que la ficha declara en la columna "Periodo" de 3_COI_CA se
        contrasta aquí con la de los pedidos de 2_NS (de donde sale el PD) y
        una diferencia queda como advertencia; los montos no se tocan."""
        if not periodos_declarados or not pedidos_por_numero:
            return
        fechas = [timezone.localdate(p.fecha_solicitud) for p in pedidos_por_numero.values()]
        desde_pd, hasta_pd = min(fechas), max(fechas)
        for texto in sorted(periodos_declarados):
            coincidencia = PATRON_PERIODO.search(texto)
            if not coincidencia:
                resumen.advertir(f'3_COI_CA: no se pudo interpretar el periodo "{texto}".')
                continue
            desde_ca, hasta_ca = (
                datetime.strptime(f, '%d/%m/%Y').date() for f in coincidencia.groups()
            )
            if (desde_ca, hasta_ca) != (desde_pd, hasta_pd):
                resumen.advertir(
                    f'3_COI_CA declara el periodo {desde_ca:%d/%m/%Y}–{hasta_ca:%d/%m/%Y}, '
                    f'pero los pedidos de 2_NS (de los que sale el PD) van del '
                    f'{desde_pd:%d/%m/%Y} al {hasta_pd:%d/%m/%Y}: CA y PD miden ventanas '
                    f'distintas. Los costos se cargaron igual, en el mes {periodo_mes:%m/%Y}.'
                )

    # ------------------------------------------------------------------
    # 3_COI_PD - Pérdidas por desabastecimiento
    # ------------------------------------------------------------------

    def _leer_coi_pd(self, hoja, resumen):
        """Lee las filas de 3_COI_PD: N.° de pedido de la ficha 2, código,
        cantidad no atendida, precio de venta y costo de compra. Las filas
        con un valor no numérico se omiten con advertencia."""
        nombre_hoja = '3_COI_PD'
        fila_encabezado = self._encontrar_fila_encabezado(hoja, nombre_hoja)
        columnas = self._mapear_columnas(
            hoja, fila_encabezado,
            alias_por_columna={
                'numero_pedido': ['n de pedido en ficha 2'],
                'codigo': ['codigo del producto'],
                'cantidad_no_atendida': ['cant no atendida', 'cantidad no atendida'],
                'precio_venta': ['precio de venta s', 'precio de venta'],
                'costo_compra': ['costo de compra s', 'costo de compra'],
            },
            requeridas=[
                'numero_pedido', 'codigo', 'cantidad_no_atendida',
                'precio_venta', 'costo_compra',
            ],
            nombre_hoja=nombre_hoja,
        )

        filas = []
        for fila in range(fila_encabezado + 1, hoja.max_row + 1):
            numero_pedido = hoja.cell(row=fila, column=columnas['numero_pedido']).value
            if numero_pedido is None:
                continue

            codigo_valor = hoja.cell(row=fila, column=columnas['codigo']).value
            cantidad_no_atendida_valor = hoja.cell(row=fila, column=columnas['cantidad_no_atendida']).value
            precio_venta_valor = hoja.cell(row=fila, column=columnas['precio_venta']).value
            costo_compra_valor = hoja.cell(row=fila, column=columnas['costo_compra']).value

            try:
                cantidad_no_atendida_excel = (
                    self._a_entero_columna(
                        resumen, nombre_hoja, fila, 'Cant. no atendida', cantidad_no_atendida_valor,
                    ) if cantidad_no_atendida_valor is not None else None
                )
                precio_venta = self._a_decimal_columna(
                    resumen, nombre_hoja, fila, 'Precio de venta (S/)', precio_venta_valor,
                )
                costo_compra = self._a_decimal_columna(
                    resumen, nombre_hoja, fila, 'Costo de compra (S/)', costo_compra_valor,
                )
            except FilaInvalida:
                resumen.filas_omitidas += 1
                continue

            filas.append({
                'fila': fila,
                'numero_pedido': numero_pedido,
                'codigo': str(codigo_valor).strip() if codigo_valor else None,
                'cantidad_no_atendida': cantidad_no_atendida_excel,
                'precio_venta': precio_venta,
                'costo_compra': costo_compra,
            })
        return filas

    def _cargar_coi_pd(self, filas_pd, pedidos_por_numero, productos_por_codigo, resumen):
        """Valida, para cada pedido no atendido por completo, que la
        cantidad no atendida coincida con la calculada en el sistema, y
        actualiza precio_venta/costo_compra del producto con los valores de
        esta hoja. El precio y el costo de la LÍNEA ya quedaron congelados
        con estos mismos valores al crearla (_cargar_ns); actualizar el
        producto solo deja el catálogo con los precios del periodo cargado y
        no cambia ninguna línea ya registrada."""
        nombre_hoja = '3_COI_PD'
        for fila_pd in filas_pd:
            fila = fila_pd['fila']
            numero_pedido = fila_pd['numero_pedido']
            codigo = fila_pd['codigo']
            precio_venta = fila_pd['precio_venta']
            costo_compra = fila_pd['costo_compra']
            cantidad_no_atendida_excel = fila_pd['cantidad_no_atendida']

            if codigo:
                producto = productos_por_codigo.get(codigo) or Producto.objects.filter(codigo=codigo).first()
                if producto is None:
                    resumen.advertir(
                        f'{nombre_hoja} fila {fila}: no se encontró el producto con '
                        f'código "{codigo}" para actualizar su precio.'
                    )
                else:
                    cambios = []
                    if producto.precio_venta != precio_venta:
                        producto.precio_venta = precio_venta
                        cambios.append('precio_venta')
                    if producto.costo_compra != costo_compra:
                        producto.costo_compra = costo_compra
                        cambios.append('costo_compra')
                    if cambios:
                        producto.save(update_fields=cambios)
                        productos_por_codigo[codigo] = producto

            pedido = pedidos_por_numero.get(numero_pedido)
            if pedido is None:
                resumen.filas_omitidas += 1
                resumen.advertir(
                    f'{nombre_hoja} fila {fila}: el pedido N.° {numero_pedido} de 2_NS '
                    f'no existe, no se pudo validar.'
                )
                continue

            detalle = pedido.detalles.first()
            cantidad_no_atendida_real = detalle.cantidad_solicitada - detalle.cantidad_atendida
            if cantidad_no_atendida_excel is not None and cantidad_no_atendida_real != cantidad_no_atendida_excel:
                resumen.advertir(
                    f'{nombre_hoja} fila {fila}: para el pedido N.° {numero_pedido} la '
                    f'cantidad no atendida del Excel ({cantidad_no_atendida_excel}) no '
                    f'coincide con la calculada en el sistema ({cantidad_no_atendida_real}).'
                )

    # ------------------------------------------------------------------
    # Reporte
    # ------------------------------------------------------------------

    def _imprimir_resumen(self, resumen):
        self.stdout.write(self.style.SUCCESS('\nResumen de la importación:'))
        self.stdout.write(f'  Productos creados:  {resumen.productos_creados}')
        self.stdout.write(f'  Pedidos creados:    {resumen.pedidos_creados}')
        self.stdout.write(f'  Movimientos creados:{resumen.movimientos_creados}')
        self.stdout.write(f'  Conteos creados:    {resumen.conteos_creados}')
        self.stdout.write(f'  Costos creados:     {resumen.costos_creados}')
        self.stdout.write(f'  Filas omitidas:     {resumen.filas_omitidas}')

        if resumen.advertencias:
            self.stdout.write(self.style.WARNING(f'\nAdvertencias ({len(resumen.advertencias)}):'))
            for advertencia in resumen.advertencias:
                self.stdout.write(f'  - {advertencia}')

        self.stdout.write(f'\nFechas no interpretadas: {len(resumen.fechas_invalidas)}')
        for hoja, fila, columna, valor in resumen.fechas_invalidas:
            self.stdout.write(f'  - {hoja} fila {fila}, columna "{columna}": valor "{valor}"')

    def _imprimir_indicadores(self, origen, meses_cargados):
        """Calcula los tres indicadores desde la base de datos (no desde el
        Excel), para verificar que la importación no deformó los datos.
        Solo del periodo que trajo el archivo (los meses en que dejó datos):
        con el pretest ya cargado, mezclar los dos periodos no verificaría
        nada.

        Usa inventario.servicios.indicadores (las mismas funciones que el dashboard),
        para que nunca haya dos cálculos distintos del mismo indicador."""
        if meses_cargados:
            desde = min(meses_cargados)
            ultimo = max(meses_cargados)
            siguiente = (ultimo.replace(day=28) + timedelta(days=4)).replace(day=1)
            hasta = siguiente - timedelta(days=1)
        else:
            desde = hasta = None
        ei = calcular_ei(desde, hasta, origen=origen)
        ns = calcular_ns(desde, hasta, origen=origen)
        coi = calcular_coi(desde, hasta, origen=origen)

        ei_valor = ei['valor'] or 0.0
        ns_valor = ns['valor'] or 0.0

        periodo = f' del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}' if desde else ''
        self.stdout.write(self.style.SUCCESS(f'\nIndicadores calculados desde la base de datos{periodo}:'))
        self.stdout.write(f'  EI (exactitud del inventario): {ei_valor:.2f}% ({ei["correctos"]}/{ei["total"]})')
        self.stdout.write(f'  NS (nivel de servicio):        {ns_valor:.2f}% ({ns["a_tiempo"]}/{ns["total"]})')
        self.stdout.write(
            f'  COI (costos operativos de inventario): S/ {coi["valor"]:.2f} '
            f'(almacenamiento S/ {coi["almacenamiento"]:.2f} + '
            f'desabastecimiento S/ {coi["desabastecimiento"]:.2f})'
        )
        self.stdout.write(
            f'    CA: {len(coi["conceptos_almacenamiento"])} concepto(s) · '
            f'PD: {coi["lineas_desabastecimiento"]} línea(s), '
            f'{coi["unidades_no_atendidas"]} unidad(es) no atendida(s), '
            f'{coi["productos_afectados"]} producto(s), al margen congelado de cada línea'
        )

        if coi['desabastecimiento'] == 0 and coi['lineas_desabastecimiento']:
            self.stdout.write(self.style.WARNING(
                '  Nota: la pérdida por desabastecimiento salió en 0 porque las líneas '
                'no atendidas no tienen precio de venta/costo de compra (3_COI_PD no '
                'trae esos pedidos y los productos no tenían precio al importar). '
                'Revisa las advertencias anteriores y, si hace falta, completa con '
                '"cargar_precios --completar-pedidos".'
            ))
