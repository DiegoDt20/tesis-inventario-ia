"""Comando de gestión: actualiza precio de venta, costo de compra y
(opcionalmente) stock mínimo de productos existentes desde un Excel.

cargar_datos crea los productos que aparecen en las fichas con precio y
costo en 0 (las fichas no traen esos datos), y un margen de 0 anula las
pérdidas por desabastecimiento del COI. Este comando completa esos campos
sin crear productos nuevos: los códigos que no existen en el catálogo solo
se reportan.
"""
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import to_excel
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import F, Q
from django.utils import timezone

from inventario.models import PedidoDetalle, Producto

from .cargar_datos import Command as CargarDatos

# Alias aceptados para cada columna, ya normalizados (minúsculas, sin
# tildes ni puntuación; ver CargarDatos._normalizar_texto). El sufijo "s"
# es lo que queda de "(S/)" tras normalizar: "Precio de venta (S/)".
# "Precio de compra" es el costo unitario de compra.
ALIAS_COLUMNAS = {
    'codigo': {'codigo', 'cod', 'codigo producto', 'codigo del producto'},
    'precio_venta': {'precio de venta', 'precio venta', 'precio', 'pv', 'precio de venta s', 'precio venta s'},
    'costo_compra': {
        'costo de compra', 'costo compra', 'costo', 'cc', 'costo de compra s',
        'precio de compra', 'precio compra', 'precio de compra s', 'precio compra s',
    },
    'stock_minimo': {'stock minimo', 'stock min', 'minimo'},
}

# Stock mínimo es opcional: si el archivo no trae la columna, ese campo no
# se toca (una lista de precios no tiene por qué traerlo).
COLUMNAS_OBLIGATORIAS = ('codigo', 'precio_venta', 'costo_compra')

ETIQUETAS_COLUMNAS = {
    'precio_venta': 'precio de venta',
    'costo_compra': 'costo de compra',
    'stock_minimo': 'stock mínimo',
}

# Filas iniciales donde se busca el encabezado (puede haber un título antes).
MAX_FILAS_BUSQUEDA_ENCABEZADO = 20

# Cuántos códigos se listan como máximo en cada sección del reporte.
MAX_CODIGOS_LISTADOS = 20


class FilaInvalida(Exception):
    """La fila se omite; el detalle ya quedó registrado como advertencia."""


@dataclass
class Resumen:
    actualizados: int = 0
    sin_cambios: int = 0
    filas_omitidas: int = 0
    celdas_fecha_convertidas: int = 0
    # Solo con --completar-pedidos (None = la opción no se usó).
    lineas_pedido_completadas: int = None
    lineas_pedido_no_atendidas_completadas: int = None
    codigos_inexistentes: list = field(default_factory=list)
    advertencias: list = field(default_factory=list)


class Command(BaseCommand):
    help = (
        'Actualiza precio de venta, costo de compra y, si el archivo trae la '
        'columna, stock mínimo de los productos existentes desde un Excel con '
        'las columnas código, precio de venta, costo (o precio) de compra y, '
        'opcional, stock mínimo. Los precios que openpyxl lea como fecha se '
        'convierten a su número de serie de Excel, con una advertencia por celda. '
        'No crea productos: reporta los códigos que no existen en el catálogo y '
        'los productos activos que siguen sin precio (precio de venta o costo de '
        'compra en 0).'
    )

    def add_arguments(self, parser):
        parser.add_argument('--archivo', required=True, help='Ruta al archivo Excel (.xlsx) de precios.')
        parser.add_argument(
            '--hoja', default=None,
            help='Nombre de la hoja a leer (por defecto, la primera del libro).',
        )
        parser.add_argument(
            '--completar-pedidos', action='store_true',
            help=(
                'Además, completa las líneas de pedido de los productos del archivo cuyo '
                'precio de venta o costo de compra guardado siga en 0 (p. ej. pedidos '
                'importados antes de conocer el precio), con los valores del archivo, y '
                'las marca como reconstruidas (precios_reconstruidos). Solo se llena el '
                'campo que está en 0; un valor ya guardado distinto de 0 no se toca.'
            ),
        )
        parser.add_argument(
            '--dry-run', action='store_true',
            help='Procesa y reporta el resultado sin escribir cambios en la base de datos.',
        )

    def handle(self, *args, **options):
        ruta = Path(options['archivo'])
        if not ruta.exists():
            raise CommandError(f'No se encontró el archivo: {ruta}')
        try:
            libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        except Exception as exc:
            raise CommandError(f'No se pudo abrir el archivo Excel: {exc}')

        nombre_hoja = options['hoja'] or libro.sheetnames[0]
        if nombre_hoja not in libro.sheetnames:
            raise CommandError(
                f'El archivo no tiene la hoja "{nombre_hoja}". Hojas disponibles: '
                f'{", ".join(libro.sheetnames)}.'
            )
        filas = list(libro[nombre_hoja].iter_rows(values_only=True))
        libro.close()

        fila_encabezado, columnas = self._mapear_encabezado(filas, nombre_hoja)
        resumen = Resumen()
        valores_por_codigo = self._leer_filas(filas, fila_encabezado, columnas, resumen)

        productos = {
            p.codigo.strip().upper(): p
            for p in Producto.objects.filter(
                codigo__in=self._variantes_codigo(valores_por_codigo)
            )
        }
        a_guardar = []
        encontrados = []
        for codigo, (fila, valores) in valores_por_codigo.items():
            producto = productos.get(codigo.upper())
            if producto is None:
                resumen.codigos_inexistentes.append(codigo)
                continue
            encontrados.append((producto, valores))
            if valores['precio_venta'] < valores['costo_compra']:
                resumen.advertencias.append(
                    f'Fila {fila}, código {codigo}: el precio de venta '
                    f'({valores["precio_venta"]}) es menor que el costo de compra '
                    f'({valores["costo_compra"]}); el margen queda negativo.'
                )
            if all(getattr(producto, campo) == valor for campo, valor in valores.items()):
                resumen.sin_cambios += 1
                continue
            for campo, valor in valores.items():
                setattr(producto, campo, valor)
            a_guardar.append(producto)
        resumen.actualizados = len(a_guardar)

        if not options['dry_run']:
            # bulk_update no dispara auto_now: actualizado_en se fija a mano.
            ahora = timezone.now()
            for producto in a_guardar:
                producto.actualizado_en = ahora

        with transaction.atomic():
            if not options['dry_run']:
                Producto.objects.bulk_update(
                    a_guardar, [c for c in ALIAS_COLUMNAS if c in columnas and c != 'codigo'] + ['actualizado_en'],
                )
            if options['completar_pedidos']:
                self._completar_pedidos(encontrados, options['dry_run'], resumen)

        self._imprimir_resumen(resumen, options['dry_run'], a_guardar, 'stock_minimo' in columnas)

    @staticmethod
    def _completar_pedidos(encontrados, dry_run, resumen):
        """Llena el precio/costo congelado de las líneas de pedido que siguen
        en 0 con los valores del archivo. Esas líneas se registraron cuando
        el producto aún no tenía precio (cargar_datos crea los productos
        nuevos en 0), y un margen de 0 congelado haría que su
        desabastecimiento no cuente nunca en el COI. Se hace con
        queryset.update() a propósito: PedidoDetalle.save() prohíbe cambiar
        un precio ya fijado, y esta es la única corrección permitida, que
        queda marcada con precios_reconstruidos=True."""
        lineas = set()
        no_atendidas = set()
        for producto, valores in encontrados:
            base = PedidoDetalle.objects.filter(producto=producto)
            por_campo = (
                ('precio_venta_unitario', valores['precio_venta']),
                ('costo_compra_unitario', valores['costo_compra']),
            )
            for campo, valor in por_campo:
                if valor <= 0:
                    continue  # completar un 0 con otro 0 no corrige nada
                qs = base.filter(**{campo: 0})
                ids = set(qs.values_list('pk', flat=True))
                if not ids:
                    continue
                lineas |= ids
                no_atendidas |= set(
                    qs.filter(cantidad_atendida__lt=F('cantidad_solicitada')).values_list('pk', flat=True)
                )
                if not dry_run:
                    qs.update(**{campo: valor, 'precios_reconstruidos': True})
        resumen.lineas_pedido_completadas = len(lineas)
        resumen.lineas_pedido_no_atendidas_completadas = len(no_atendidas)

    # ------------------------------------------------------------------
    # Lectura del Excel
    # ------------------------------------------------------------------

    def _mapear_encabezado(self, filas, nombre_hoja):
        """Busca la primera fila que tenga las columnas obligatorias y
        devuelve (índice de fila, {campo: índice de columna}); stock_minimo
        solo aparece en el dict si el archivo trae esa columna."""
        for i, fila in enumerate(filas[:MAX_FILAS_BUSQUEDA_ENCABEZADO]):
            encontrados = {}
            for j, celda in enumerate(fila):
                texto = CargarDatos._normalizar_texto(celda)
                for campo, alias in ALIAS_COLUMNAS.items():
                    if campo not in encontrados and texto in alias:
                        encontrados[campo] = j
            if all(c in encontrados for c in COLUMNAS_OBLIGATORIAS):
                return i, encontrados
            if 'codigo' in encontrados:
                faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in encontrados]
                raise CommandError(
                    f'{nombre_hoja}: en la fila de encabezado (fila {i + 1}) faltan las '
                    f'columnas {faltantes}. Se esperan: código, precio de venta y '
                    'costo de compra (stock mínimo es opcional).'
                )
        raise CommandError(
            f'{nombre_hoja}: no se encontró la fila de encabezado con la columna '
            f'"código" en las primeras {MAX_FILAS_BUSQUEDA_ENCABEZADO} filas.'
        )

    def _leer_filas(self, filas, fila_encabezado, columnas, resumen):
        """Devuelve {código: (número de fila en Excel, valores)}. Omite, con
        advertencia, las filas con algún valor vacío, no numérico o
        negativo. Si un código se repite, gana la última fila."""
        valores_por_codigo = {}
        for i in range(fila_encabezado + 1, len(filas)):
            fila = filas[i]
            numero_fila = i + 1
            celda_codigo = fila[columnas['codigo']] if columnas['codigo'] < len(fila) else None
            codigo = self._a_codigo(celda_codigo)
            if not codigo:
                if any(c not in (None, '') for c in fila):
                    resumen.filas_omitidas += 1
                    resumen.advertencias.append(f'Fila {numero_fila}: sin código, se omitió.')
                continue
            try:
                valores = {
                    'precio_venta': self._a_monto(fila, columnas, 'precio_venta', numero_fila, resumen),
                    'costo_compra': self._a_monto(fila, columnas, 'costo_compra', numero_fila, resumen),
                }
                if 'stock_minimo' in columnas:
                    valores['stock_minimo'] = self._a_entero(fila, columnas, 'stock_minimo', numero_fila, resumen)
            except FilaInvalida:
                resumen.filas_omitidas += 1
                continue
            if codigo in valores_por_codigo:
                resumen.advertencias.append(
                    f'Fila {numero_fila}: el código {codigo} ya apareció en la fila '
                    f'{valores_por_codigo[codigo][0]}; se usan los valores de la fila {numero_fila}.'
                )
            valores_por_codigo[codigo] = (numero_fila, valores)
        return valores_por_codigo

    @staticmethod
    def _a_codigo(valor):
        """Excel convierte los códigos puramente numéricos en número
        (123 -> 123.0); se devuelven como texto sin el ".0"."""
        if valor is None:
            return ''
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        return str(valor).strip()

    @staticmethod
    def _variantes_codigo(valores_por_codigo):
        """El código se compara sin distinguir mayúsculas; como filtrar con
        __iexact en bloque no es posible, se buscan las variantes exacta,
        en mayúsculas y en minúsculas."""
        variantes = set()
        for codigo in valores_por_codigo:
            variantes.update({codigo, codigo.upper(), codigo.lower()})
        return variantes

    @staticmethod
    def _leer_numero(fila, columnas, campo, numero_fila, resumen):
        indice = columnas[campo]
        valor = fila[indice] if indice < len(fila) else None
        etiqueta = ETIQUETAS_COLUMNAS[campo]
        if valor is None or (isinstance(valor, str) and not valor.strip()):
            resumen.advertencias.append(f'Fila {numero_fila}: falta el {etiqueta}, se omitió la fila.')
            raise FilaInvalida
        if isinstance(valor, (datetime, date, time)):
            valor = Command._fecha_a_numero(valor, numero_fila, etiqueta, resumen)
        elif isinstance(valor, bool):
            valor = None
        else:
            try:
                valor = Decimal(str(valor).strip())
            except InvalidOperation:
                valor = None
        if valor is None or not valor.is_finite():
            resumen.advertencias.append(
                f'Fila {numero_fila}, columna "{etiqueta}": se esperaba un número y se '
                f'encontró "{fila[indice]}", se omitió la fila.'
            )
            raise FilaInvalida
        if valor < 0:
            resumen.advertencias.append(
                f'Fila {numero_fila}, columna "{etiqueta}": el valor {valor} es negativo, '
                'se omitió la fila.'
            )
            raise FilaInvalida
        return valor

    @staticmethod
    def _fecha_a_numero(valor, numero_fila, etiqueta, resumen):
        """Una celda numérica que openpyxl entregó como fecha/hora vuelve a
        su número de serie de Excel. Pasa, por ejemplo, con el formato de
        moneda "S/ #,##0.00": openpyxl toma la "S" de "S/" como el código de
        segundos de un formato de hora y convierte el número en fecha
        (45 -> 1900-02-14, 7,5 -> 1900-01-07 12:00).

        Se usa to_excel y no "días desde el 31/12/1899": Excel cuenta un
        29/02/1900 que no existió, así que desde el serial 61 esa resta da
        un día menos (un precio de 120 saldría 119)."""
        if isinstance(valor, time):
            numero = (valor.hour * 3600 + valor.minute * 60 + valor.second + valor.microsecond / 1e6) / 86400
            leido = f'la hora {valor:%H:%M:%S}'
        else:
            if not isinstance(valor, datetime):
                valor = datetime.combine(valor, time())
            numero = to_excel(valor)
            leido = f'la fecha {valor:%Y-%m-%d %H:%M}'

        # El serial se guarda como float en el .xlsx: el redondeo corrige
        # restos binarios de la conversión (7.499999999 -> 7.5).
        convertido = Decimal(str(round(numero, 6))).normalize()
        resumen.advertencias.append(
            f'Fila {numero_fila}, columna "{etiqueta}": se leyó {leido}, '
            f'se convirtió al número {convertido:f}.'
        )
        resumen.celdas_fecha_convertidas += 1
        return convertido

    def _a_monto(self, fila, columnas, campo, numero_fila, resumen):
        valor = self._leer_numero(fila, columnas, campo, numero_fila, resumen)
        if valor >= Decimal('1e10'):
            # max_digits=12, decimal_places=2: a partir de 10^10 no cabe.
            resumen.advertencias.append(
                f'Fila {numero_fila}, columna "{ETIQUETAS_COLUMNAS[campo]}": el monto '
                f'{valor} excede el máximo permitido, se omitió la fila.'
            )
            raise FilaInvalida
        return valor.quantize(Decimal('0.01'))

    def _a_entero(self, fila, columnas, campo, numero_fila, resumen):
        valor = self._leer_numero(fila, columnas, campo, numero_fila, resumen)
        if valor != valor.to_integral_value():
            resumen.advertencias.append(
                f'Fila {numero_fila}, columna "{ETIQUETAS_COLUMNAS[campo]}": se esperaba '
                f'un número entero y se encontró {valor}, se omitió la fila.'
            )
            raise FilaInvalida
        return int(valor)

    # ------------------------------------------------------------------
    # Reporte
    # ------------------------------------------------------------------

    def _imprimir_resumen(self, resumen, dry_run, actualizados, trae_stock_minimo):
        if dry_run:
            self.stdout.write(self.style.WARNING('Modo --dry-run: no se escribió ningún cambio.'))

        verbo = 'Se actualizarían' if dry_run else 'Productos actualizados'
        self.stdout.write(self.style.SUCCESS(f'{verbo}: {resumen.actualizados}'))
        self.stdout.write(f'Productos sin cambios (ya tenían esos valores): {resumen.sin_cambios}')
        self.stdout.write(f'Filas omitidas por datos inválidos: {resumen.filas_omitidas}')
        if not trae_stock_minimo:
            self.stdout.write('El archivo no trae stock mínimo: ese campo no se modificó.')
        if resumen.celdas_fecha_convertidas:
            self.stdout.write(self.style.WARNING(
                f'Celdas que llegaron como fecha y se convirtieron a número: '
                f'{resumen.celdas_fecha_convertidas} (detalle en las advertencias).'
            ))

        if resumen.lineas_pedido_completadas is not None:
            verbo_lineas = 'que se completarían' if dry_run else 'completadas'
            self.stdout.write(
                f'Líneas de pedido con precio o costo en 0 {verbo_lineas} (marcadas como '
                f'reconstruidas): {resumen.lineas_pedido_completadas}, de ellas no atendidas '
                f'por completo, que suman al COI: {resumen.lineas_pedido_no_atendidas_completadas}'
            )
        self.stdout.write(f'Códigos del archivo que no existen en el catálogo: {len(resumen.codigos_inexistentes)}')
        self._listar(resumen.codigos_inexistentes)

        # Los productos actualizados se evalúan con sus valores nuevos en
        # memoria: en --dry-run la base todavía tiene los anteriores.
        sin_precio_en_base = Producto.objects.filter(activo=True).filter(
            Q(precio_venta=0) | Q(costo_compra=0)
        ).exclude(pk__in=[p.pk for p in actualizados])
        codigos_sin_precio = sorted(
            list(sin_precio_en_base.values_list('codigo', flat=True))
            + [p.codigo for p in actualizados if p.activo and (p.precio_venta == 0 or p.costo_compra == 0)]
        )
        self.stdout.write(
            'Productos activos del catálogo sin precio (precio de venta o costo de '
            f'compra en 0): {len(codigos_sin_precio)}'
        )
        self._listar(codigos_sin_precio)

        if resumen.advertencias:
            self.stdout.write(self.style.WARNING(f'\nAdvertencias ({len(resumen.advertencias)}):'))
            for texto in resumen.advertencias:
                self.stdout.write(f'  - {texto}')

    def _listar(self, codigos):
        if not codigos:
            return
        muestra = ', '.join(codigos[:MAX_CODIGOS_LISTADOS])
        resto = len(codigos) - MAX_CODIGOS_LISTADOS
        self.stdout.write(f'  {muestra}' + (f' ... y {resto} más' if resto > 0 else ''))
