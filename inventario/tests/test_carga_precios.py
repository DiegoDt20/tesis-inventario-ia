"""Tests del comando cargar_precios: actualización de precio, costo y
stock mínimo desde Excel, y el reporte de códigos inexistentes y productos
sin precio."""
import tempfile
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from ..models import Producto


class CargarPreciosTests(TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        for codigo in ('PIN-001', 'PIN-002', 'PIN-003'):
            Producto.objects.create(codigo=codigo, nombre=codigo, precio_venta=0, costo_compra=0)

    def _excel(self, filas, encabezado=('Código', 'Precio de venta', 'Costo de compra', 'Stock mínimo')):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Lista de precios septiembre'])
        hoja.append(list(encabezado))
        for fila in filas:
            hoja.append(list(fila))
        ruta = Path(self.dir.name) / 'precios.xlsx'
        libro.save(ruta)
        return ruta

    def _correr(self, ruta, **opciones):
        salida = StringIO()
        call_command('cargar_precios', archivo=str(ruta), stdout=salida, **opciones)
        return salida.getvalue()

    def test_actualiza_y_reporta(self):
        ruta = self._excel([
            ('PIN-001', 55.5, 30, 4),
            ('pin-002', '40.00', '25.10', 2),
            ('PIN-999', 10, 5, 1),
        ])
        salida = self._correr(ruta)

        p1 = Producto.objects.get(codigo='PIN-001')
        self.assertEqual((p1.precio_venta, p1.costo_compra, p1.stock_minimo), (Decimal('55.50'), Decimal('30.00'), 4))
        p2 = Producto.objects.get(codigo='PIN-002')
        self.assertEqual((p2.precio_venta, p2.costo_compra, p2.stock_minimo), (Decimal('40.00'), Decimal('25.10'), 2))
        self.assertIn('Productos actualizados: 2', salida)
        self.assertIn('no existen en el catálogo: 1', salida)
        self.assertIn('PIN-999', salida)
        self.assertIn('sin precio (precio de venta o costo de compra en 0): 1', salida)
        self.assertIn('PIN-003', salida)

    def test_filas_invalidas_se_omiten(self):
        ruta = self._excel([
            ('PIN-001', 'abc', 30, 4),
            ('PIN-002', 40, -1, 2),
            ('PIN-003', 40, 20, 2.5),
            ('PIN-003', 40, 20, None),
        ])
        salida = self._correr(ruta)
        self.assertIn('Productos actualizados: 0', salida)
        self.assertIn('Filas omitidas por datos inválidos: 4', salida)
        self.assertFalse(Producto.objects.exclude(precio_venta=0).exists())

    def test_dry_run_no_escribe_pero_reporta(self):
        ruta = self._excel([('PIN-001', 50, 30, 4), ('PIN-002', 50, 30, 4), ('PIN-003', 50, 30, 4)])
        salida = self._correr(ruta, dry_run=True)
        self.assertIn('Se actualizarían: 3', salida)
        self.assertIn('en 0): 0', salida)
        self.assertFalse(Producto.objects.exclude(precio_venta=0).exists())

    def test_sin_cambios_no_cuenta_como_actualizado(self):
        Producto.objects.filter(codigo='PIN-001').update(precio_venta=50, costo_compra=30, stock_minimo=4)
        salida = self._correr(self._excel([('PIN-001', 50, 30, 4)]))
        self.assertIn('Productos actualizados: 0', salida)
        self.assertIn('sin cambios (ya tenían esos valores): 1', salida)

    def test_precio_menor_al_costo_advierte(self):
        salida = self._correr(self._excel([('PIN-001', 20, 30, 4)]))
        self.assertIn('Productos actualizados: 1', salida)
        self.assertIn('margen queda negativo', salida)

    def test_falta_columna_obligatoria(self):
        ruta = self._excel([('PIN-001', 50, 4)], encabezado=('Código', 'Precio de venta', 'Stock mínimo'))
        with self.assertRaisesMessage(CommandError, 'costo_compra'):
            self._correr(ruta)

    def test_marca_actualizado_en(self):
        antes = Producto.objects.get(codigo='PIN-001').actualizado_en
        self._correr(self._excel([('PIN-001', 50, 30, 4)]))
        self.assertGreater(Producto.objects.get(codigo='PIN-001').actualizado_en, antes)

    def test_precios_leidos_como_fecha_se_convierten(self):
        # Formato de moneda "S/ #,##0.00": openpyxl confunde la "S" con el
        # código de segundos y entrega los números como fechas al leer.
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Código', 'Producto', 'Precio de venta (S/)', 'Precio de compra (S/)'])
        hoja.append(['PIN-001', 'Innova', 45, 7.5])
        hoja.append(['PIN-002', 'Latex', 120, 100])  # desde el serial 61 cuenta el 29/02/1900 de Excel
        for fila in hoja.iter_rows(min_row=2, min_col=3, max_col=4):
            for celda in fila:
                celda.number_format = 'S/ #,##0.00'
        ruta = Path(self.dir.name) / 'precios_fecha.xlsx'
        libro.save(ruta)
        Producto.objects.filter(codigo='PIN-001').update(stock_minimo=9)

        salida = self._correr(ruta)

        p1 = Producto.objects.get(codigo='PIN-001')
        self.assertEqual((p1.precio_venta, p1.costo_compra), (Decimal('45.00'), Decimal('7.50')))
        self.assertEqual(p1.stock_minimo, 9)  # el archivo no trae stock mínimo: no se toca
        p2 = Producto.objects.get(codigo='PIN-002')
        self.assertEqual((p2.precio_venta, p2.costo_compra), (Decimal('120.00'), Decimal('100.00')))
        self.assertIn('Celdas que llegaron como fecha y se convirtieron a número: 4', salida)
        self.assertIn('Fila 2, columna "precio de venta": se leyó la fecha 1900-02-14 00:00, se convirtió al número 45.', salida)
        self.assertIn('Fila 2, columna "costo de compra": se leyó la fecha 1900-01-07 12:00, se convirtió al número 7.5.', salida)
        self.assertIn('se convirtió al número 100.', salida)
        self.assertIn('El archivo no trae stock mínimo', salida)
