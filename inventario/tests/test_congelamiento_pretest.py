"""El pretest queda congelado al importar el postest.

Es el riesgo más grande de la comparación pretest/postest: si importar los
datos del periodo posterior o actualizar precios cambiara un decimal de EI,
NS o COI de agosto, la comparación dejaría de ser válida. Estas pruebas
cargan un pretest y un postest con el mismo comando que se usa con las
fichas reales (cargar_datos, origen "real") y actualizan precios con
cargar_precios, como se hará de verdad.
"""
import json
import tempfile
from datetime import date, datetime
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.contrib.auth.models import User
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from ..models import (
    ConteoDetalle,
    CostoAlmacenamiento,
    Movimiento,
    Pedido,
    PedidoDetalle,
    Producto,
)
from ..servicios.indicadores import calcular_coi, calcular_ei, calcular_ns

AGOSTO = (date(2026, 8, 1), date(2026, 8, 31))
SETIEMBRE = (date(2026, 9, 1), date(2026, 9, 30))


def _ficha(ruta, conteo, pedidos, costos, perdidas):
    """Arma una ficha de registro con la estructura de las reales
    (datos/fichas_pretest_agosto_2026.xlsx): mismas hojas y encabezados.

    conteo:   (fecha_corte, [(codigo, descripcion, kardex, fisico), ...])
    pedidos:  [(n, fecha, cliente, codigo, descripcion, solicitada, atendida,
                requerida, atencion, a_tiempo), ...]
    costos:   (periodo_texto, [(componente, monto), ...])
    perdidas: [(n_pedido, codigo, no_atendida, precio_venta, costo_compra), ...]
    """
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    ei = libro.create_sheet('1_EI')
    for _ in range(7):
        ei.append([])
    ei.append([
        'N.°', 'Fecha de corte', 'Código', 'Producto y presentación',
        'Stock según kardex', 'Stock según conteo físico', 'Diferencia',
        '¿Registro correcto? (1/0)',
    ])
    fecha_corte, filas_conteo = conteo
    for n, (codigo, descripcion, kardex, fisico) in enumerate(filas_conteo, start=1):
        ei.append([n, fecha_corte, codigo, descripcion, kardex, fisico, fisico - kardex, int(kardex == fisico)])

    ns = libro.create_sheet('2_NS')
    for _ in range(7):
        ns.append([])
    ns.append([
        'N.°', 'Fecha del pedido', 'Cliente (código)', 'Código del producto',
        'Producto y presentación', 'Cant. solicitada', 'Cant. atendida',
        'Fecha requerida', 'Fecha de atención', 'Motivo', '¿A tiempo? (1/0)',
    ])
    for n, fecha, cliente, codigo, descripcion, sol, ate, req, aten, a_tiempo in pedidos:
        ns.append([n, fecha, cliente, codigo, descripcion, sol, ate, req, aten, None, a_tiempo])

    ca = libro.create_sheet('3_COI_CA')
    periodo, filas_costos = costos
    ca.append(['FICHA N.° 3 — BLOQUE A: COSTOS DE ALMACENAMIENTO (CA)'])
    ca.append([f'Costos registrados por concepto para el periodo {periodo}'])
    ca.append([])
    ca.append(['Componente', 'Detalle del cálculo', 'Periodo', 'Monto (S/)'])
    for componente, monto in filas_costos:
        ca.append([componente, 'detalle', periodo, monto])
    ca.append(['TOTAL DE COSTOS DE ALMACENAMIENTO (CA)', None, None, sum(m for _, m in filas_costos)])

    pd = libro.create_sheet('3_COI_PD')
    for _ in range(3):
        pd.append([])
    pd.append([
        'N.°', 'N.° de pedido en Ficha 2', 'Código del producto', 'Producto y presentación',
        'Cant. no atendida', 'Precio de venta (S/)', 'Costo de compra (S/)',
        'Margen unitario (S/)', 'Pérdida (S/)',
    ])
    for n, (n_pedido, codigo, no_atendida, pv, cc) in enumerate(perdidas, start=1):
        pd.append([n, n_pedido, codigo, codigo, no_atendida, pv, cc, pv - cc, no_atendida * (pv - cc)])
    pd.append(['TOTAL DE PÉRDIDAS POR DESABASTECIMIENTO (PD)'])

    libro.save(ruta)
    return ruta


PRODUCTOS = [
    ('PIN-001', 'CPP Pato — Blanco — galón'),
    ('PIN-002', 'Innova — Celeste — balde 4 galones'),
    ('PIN-003', 'Tumi — Brocha — 3 pulgadas'),
    ('PIN-004', 'Vencedor — Crema — galón'),
]


def _pretest_agosto(ruta):
    """Pretest: conteo al 24/08, pedidos del 01/08 al 31/08 y CA con la
    ventana 25/07–24/08, como la ficha real."""
    d = lambda dia: date(2026, 8, dia)  # noqa: E731
    return _ficha(
        ruta,
        conteo=(d(24), [
            ('PIN-001', PRODUCTOS[0][1], 12, 10),
            ('PIN-002', PRODUCTOS[1][1], 6, 6),
            ('PIN-003', PRODUCTOS[2][1], 14, 16),
            ('PIN-004', PRODUCTOS[3][1], 15, 15),
        ]),
        pedidos=[
            (1, d(1), 'C-001', 'PIN-001', PRODUCTOS[0][1], 8, 8, d(1), d(1), 1),
            (2, d(5), 'C-002', 'PIN-001', PRODUCTOS[0][1], 12, 10, d(5), d(5), 0),
            (3, d(12), 'C-003', 'PIN-002', PRODUCTOS[1][1], 5, 0, d(13), None, 0),
            (4, d(20), 'C-004', 'PIN-003', PRODUCTOS[2][1], 2, 2, d(21), d(22), 0),
            (5, d(31), 'C-005', 'PIN-004', PRODUCTOS[3][1], 3, 3, d(31), d(31), 1),
        ],
        costos=('25/07/2026–24/08/2026', [
            ('Alquiler imputado al almacén', 550),
            ('Energía eléctrica imputada', 50),
            ('Personal de almacén', 300),
            ('Mermas del periodo', 150),
        ]),
        # PD = 2 × (34 − 25.5) + 5 × (120 − 90) = 17 + 150 = 167
        perdidas=[(2, 'PIN-001', 2, 34, 25.5), (3, 'PIN-002', 5, 120, 90)],
    )


def _postest_setiembre(ruta, primer_pedido=date(2026, 9, 2)):
    """Postest: los mismos productos con OTROS precios en 3_COI_PD (la hoja
    de precios del postest actualiza el catálogo), conteo al 25/09 y CA con
    la ventana 26/08–25/09, como el ejemplo de postest."""
    d = lambda dia: date(2026, 9, dia)  # noqa: E731
    return _ficha(
        ruta,
        conteo=(d(25), [
            ('PIN-001', PRODUCTOS[0][1], 10, 10),
            ('PIN-002', PRODUCTOS[1][1], 7, 6),
            ('PIN-003', PRODUCTOS[2][1], 16, 16),
            ('PIN-004', PRODUCTOS[3][1], 15, 15),
        ]),
        pedidos=[
            (1, primer_pedido, 'C-001', 'PIN-001', PRODUCTOS[0][1], 10, 6, primer_pedido, primer_pedido, 0),
            (2, d(9), 'C-006', 'PIN-002', PRODUCTOS[1][1], 4, 1, d(9), d(9), 0),
            (3, d(15), 'C-007', 'PIN-004', PRODUCTOS[3][1], 6, 6, d(16), d(16), 1),
        ],
        costos=('26/08/2026–25/09/2026', [
            ('Alquiler imputado al almacén', 535),
            ('Energía eléctrica imputada', 48),
        ]),
        perdidas=[(1, 'PIN-001', 4, 40, 28), (2, 'PIN-002', 3, 150, 100)],
    )


def _precios(ruta, filas):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(['Código', 'Precio de venta', 'Costo de compra'])
    for fila in filas:
        hoja.append(list(fila))
    libro.save(ruta)
    return ruta


def _indicadores(desde, hasta):
    return {
        'ei': calcular_ei(desde, hasta, origen='real'),
        'ns': calcular_ns(desde, hasta, origen='real'),
        'coi': calcular_coi(desde, hasta, origen='real'),
    }


class _ConFichas(TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        self.ruta = lambda nombre: Path(self.dir.name) / nombre  # noqa: E731

    def _cargar(self, ruta, **opciones):
        salida = StringIO()
        call_command(
            'cargar_datos', archivo=str(ruta), origen='real', stdout=salida, stderr=StringIO(), **opciones,
        )
        return salida.getvalue()

    def _cargar_pretest(self):
        return self._cargar(_pretest_agosto(self.ruta('pretest.xlsx')))


class CongelamientoPretestTests(_ConFichas):
    def test_postest_y_precios_nuevos_no_cambian_el_pretest(self):
        # 1. Indicadores de agosto, recién cargado el pretest.
        self._cargar_pretest()
        antes = _indicadores(*AGOSTO)
        # El pretest no es trivial: hay CA, PD y registros incorrectos.
        self.assertEqual((antes['ei']['correctos'], antes['ei']['total']), (2, 4))
        self.assertEqual((antes['ns']['a_tiempo'], antes['ns']['total']), (2, 5))
        self.assertEqual(antes['coi']['almacenamiento'], Decimal('1050.00'))
        self.assertEqual(antes['coi']['desabastecimiento'], Decimal('167.00'))
        self.assertEqual(antes['coi']['valor'], Decimal('1217.00'))

        # 2. Pedidos, movimientos y conteos de setiembre (con la misma
        # carga que se usará para el postest real, que además pone en el
        # catálogo los precios de su hoja 3_COI_PD) y un movimiento
        # operativo registrado en el sistema.
        self._cargar(_postest_setiembre(self.ruta('postest.xlsx')))
        Movimiento.objects.create(
            producto=Producto.objects.get(codigo='PIN-001'), tipo=Movimiento.Tipo.INGRESO,
            cantidad=20, fecha=timezone.make_aware(datetime(2026, 9, 26, 10)),
            usuario=User.objects.create_user('operador'), origen='real',
        )

        # 3. Precio y costo nuevos para los productos del PD de agosto, con
        # cargar_precios (también con --completar-pedidos, que sí escribe
        # en líneas de pedido) y a mano.
        call_command(
            'cargar_precios', archivo=str(_precios(self.ruta('precios.xlsx'), [
                ('PIN-001', 99.90, 10.10), ('PIN-002', 300, 1), ('PIN-003', 7, 3),
            ])),
            completar_pedidos=True, stdout=StringIO(),
        )
        pin002 = Producto.objects.get(codigo='PIN-002')
        pin002.precio_venta, pin002.costo_compra = Decimal('500.00'), Decimal('0.50')
        pin002.save()

        # 4 y 5. Agosto da exactamente lo mismo, decimal por decimal.
        despues = _indicadores(*AGOSTO)
        self.assertEqual(despues['ei'], antes['ei'])
        self.assertEqual(despues['ns'], antes['ns'])
        for clave in ('valor', 'almacenamiento', 'desabastecimiento'):
            self.assertEqual(str(despues['coi'][clave]), str(antes['coi'][clave]), clave)
        self.assertEqual(despues['coi'], antes['coi'])

        # Y setiembre se mide aparte, con los precios de su propia ficha.
        postest = _indicadores(*SETIEMBRE)
        self.assertEqual((postest['ei']['correctos'], postest['ei']['total']), (3, 4))
        self.assertEqual((postest['ns']['a_tiempo'], postest['ns']['total']), (1, 3))
        # CA 583 + PD (4 × 12 + 3 × 50 = 198)
        self.assertEqual(postest['coi']['valor'], Decimal('781.00'))

    def test_el_pd_usa_el_precio_de_la_ficha_aunque_el_catalogo_tenga_otro(self):
        # El catálogo ya trae un precio distinto al de la hoja 3_COI_PD
        # (como pasará en el postest, con los productos del pretest).
        Producto.objects.create(codigo='PIN-001', nombre='CPP Pato', precio_venta=70, costo_compra=20, origen='real')
        self._cargar_pretest()

        linea = PedidoDetalle.objects.get(producto__codigo='PIN-001', cantidad_atendida=10)
        self.assertEqual((linea.precio_venta_unitario, linea.costo_compra_unitario), (Decimal('34.00'), Decimal('25.50')))
        self.assertFalse(linea.precios_reconstruidos)
        # Y el COI de agosto se calcula con el margen de la ficha (8.50), no
        # con el del catálogo previo (50.00).
        self.assertEqual(calcular_coi(*AGOSTO, origen='real')['desabastecimiento'], Decimal('167.00'))


class ResguardosDePeriodoTests(_ConFichas):
    def _conteos(self):
        return (Pedido.objects.count(), ConteoDetalle.objects.count(), CostoAlmacenamiento.objects.count(),
                Movimiento.objects.count())

    def test_rechaza_un_postest_con_pedidos_en_el_mes_del_pretest(self):
        self._cargar_pretest()
        antes, conteos = _indicadores(*AGOSTO), self._conteos()

        ficha = _postest_setiembre(self.ruta('postest.xlsx'), primer_pedido=date(2026, 8, 28))
        with self.assertRaisesMessage(CommandError, '08/2026'):
            self._cargar(ficha)

        self.assertEqual(self._conteos(), conteos)
        self.assertEqual(_indicadores(*AGOSTO), antes)

    def test_rechaza_cargar_dos_veces_el_mismo_periodo(self):
        self._cargar_pretest()
        conteos = self._conteos()
        with self.assertRaisesMessage(CommandError, 'ya hay datos de origen "real"'):
            self._cargar_pretest()
        self.assertEqual(self._conteos(), conteos)

    def test_limpiar_no_puede_borrar_el_pretest(self):
        self._cargar_pretest()
        antes, conteos = _indicadores(*AGOSTO), self._conteos()

        with self.assertRaisesMessage(CommandError, '--limpiar borraría'):
            self._cargar(_postest_setiembre(self.ruta('postest.xlsx')), limpiar=True)

        self.assertEqual(self._conteos(), conteos)
        self.assertEqual(_indicadores(*AGOSTO), antes)

    def test_limpiar_si_puede_recargar_el_mismo_periodo(self):
        self._cargar_pretest()
        antes = _indicadores(*AGOSTO)
        self._cargar(_pretest_agosto(self.ruta('pretest.xlsx')), limpiar=True)
        self.assertEqual(_indicadores(*AGOSTO), antes)
        self.assertEqual(Pedido.objects.count(), 5)

    def test_el_dry_run_tambien_avisa(self):
        self._cargar_pretest()
        with self.assertRaises(CommandError):
            self._cargar(_postest_setiembre(self.ruta('p.xlsx'), primer_pedido=date(2026, 8, 30)), dry_run=True)

    def test_el_reporte_de_la_carga_es_solo_del_periodo_cargado(self):
        self._cargar_pretest()
        salida = self._cargar(_postest_setiembre(self.ruta('postest.xlsx')))
        self.assertIn('del 01/09/2026 al 30/09/2026', salida)
        self.assertIn('S/ 781.00', salida)
        self.assertIn('PD: 2 línea(s), 7 unidad(es)', salida)


class FichaNSyCATests(_ConFichas):
    def test_advierte_si_la_ventana_del_ca_no_es_la_de_los_pedidos(self):
        salida = self._cargar_pretest()
        self.assertIn(
            '3_COI_CA declara el periodo 25/07/2026–24/08/2026, pero los pedidos de 2_NS '
            '(de los que sale el PD) van del 01/08/2026 al 31/08/2026', salida,
        )
        # No cambia ningún monto: el CA se cargó completo en agosto.
        self.assertEqual(calcular_coi(*AGOSTO, origen='real')['almacenamiento'], Decimal('1050.00'))

    def test_guarda_la_fecha_requerida(self):
        self._cargar_pretest()
        linea = PedidoDetalle.objects.get(pedido__cliente='C-003')
        self.assertEqual(linea.fecha_requerida, date(2026, 8, 13))

    def test_a_tiempo_vacio_se_calcula_con_las_fechas(self):
        ruta = _pretest_agosto(self.ruta('pretest.xlsx'))
        libro = openpyxl.load_workbook(ruta)
        hoja = libro['2_NS']
        for fila in range(9, 14):  # fórmula sin valor calculado
            hoja.cell(row=fila, column=11).value = None
        libro.save(ruta)

        salida = self._cargar(ruta)
        # Mismo NS que con la columna llena: 2 de 5.
        ns = calcular_ns(*AGOSTO, origen='real')
        self.assertEqual((ns['a_tiempo'], ns['total']), (2, 5))
        self.assertIn('"¿A tiempo?" vacío; se calculó', salida)

    def test_advierte_si_a_tiempo_contradice_las_fechas(self):
        ruta = _pretest_agosto(self.ruta('pretest.xlsx'))
        libro = openpyxl.load_workbook(ruta)
        libro['2_NS'].cell(row=12, column=11).value = 1  # pedido 4: atendido 22/08, requerido 21/08
        libro.save(ruta)

        salida = self._cargar(ruta)
        self.assertIn('2_NS fila 12: "¿A tiempo?" dice 1, pero con la fecha requerida', salida)
        # Manda la ficha.
        self.assertTrue(PedidoDetalle.objects.get(pedido__cliente='C-004').atendido_a_tiempo)


class ReporteIndicadoresTests(_ConFichas):
    def _reporte(self, **opciones):
        salida = StringIO()
        call_command(
            'reporte_indicadores', desde=AGOSTO[0], hasta=AGOSTO[1], stdout=salida, **opciones,
        )
        return salida.getvalue()

    def test_desglose_y_detalle(self):
        self._cargar_pretest()
        salida = self._reporte(detalle=True)
        self.assertIn('COI = CA + PD = S/ 1217.00', salida)
        self.assertIn('→ S/ 1050.00', salida)
        self.assertIn('→ S/ 167.00', salida)
        self.assertIn('2 línea(s), 7 unidad(es) no atendida(s), 2 producto(s)', salida)
        self.assertIn('05/08/2026  C-002     PIN-001            2     34.00     25.50     8.50      17.00  no', salida)

    def test_guardar_y_comparar(self):
        self._cargar_pretest()
        foto = self.ruta('pretest.json')
        self._reporte(guardar=str(foto))
        self.assertEqual(json.loads(foto.read_text())['coi']['valor'], '1217.00')

        self._cargar(_postest_setiembre(self.ruta('postest.xlsx')))
        self.assertIn('los valores son idénticos', self._reporte(comparar=str(foto)))

        # Un costo agregado a agosto sí se detecta.
        CostoAlmacenamiento.objects.create(
            periodo_mes=date(2026, 8, 1), concepto='Extra', monto=Decimal('0.01'), origen='real',
        )
        with self.assertRaisesMessage(CommandError, "coi.valor: '1217.00' → '1217.01'"):
            self._reporte(comparar=str(foto))


class DesgloseEnPantallaTests(_ConFichas):
    def test_reportes_muestra_el_desglose_del_coi(self):
        self._cargar_pretest()
        self.client.force_login(User.objects.create_superuser('admin', 'a@a.pe', 'x'))
        respuesta = self.client.get(
            reverse('inventario:reportes'), {'fecha_inicio': '2026-08-01', 'fecha_fin': '2026-08-31'},
        )
        self.assertContains(respuesta, 'Desglose del COI')
        self.assertContains(respuesta, '2 líneas no atendidas por completo')
        self.assertContains(respuesta, '7 unidades')
